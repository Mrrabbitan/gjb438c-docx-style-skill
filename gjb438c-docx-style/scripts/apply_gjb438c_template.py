#!/usr/bin/env python3
"""Apply a document-type-aware GJB 438C DOCX template to structured content.

SRS uses the dedicated P-09 three-section profile and Appendix J outline.
The overall technical solution retains its original two-section profile.
"""

from __future__ import annotations

import argparse
import json
import re
import os
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

from audit_gjb438c_docx import AuditResult, audit_page_numbering
from srs_model import (CLAUSE_TITLES, normalize_srs_content, require_valid_srs_content, validate_srs_content)


SKILL_DIR = Path(__file__).resolve().parents[1]
PROFILE_FILE = SKILL_DIR / "references" / "document-profiles.json"
DOCUMENT_PROFILES = json.loads(PROFILE_FILE.read_text(encoding="utf-8"))
DEFAULT_TEMPLATE = SKILL_DIR / DOCUMENT_PROFILES["overall-technical-solution"]["template"]
SRS_TEMPLATE = SKILL_DIR / DOCUMENT_PROFILES["SRS-P09"]["template"]

HEADING_STYLE_BY_LEVEL = {
    1: "Heading 1",
    2: "Heading 2",
    3: "Heading 3",
    4: "Heading 4",
    5: "Heading 5",
}
HEADING_LEVEL_BY_STYLE = {name: level for level, name in HEADING_STYLE_BY_LEVEL.items()}

TABLE_HEADER_STYLE = "145表头"
TABLE_BODY_STYLE = "145表正文"
FIGURE_STYLE = "145图样式"
GENERATED_TABLE_STYLE = "Table Grid1"

HEADING_NUMBER_RE = re.compile(r"^\d+(?:\.\d+)*[\s　]*")

SRS_REQUIRED_H1 = ["范围", "引用文档", "需求", "合格性规定", "需求可追踪性", "注释"]
SRS_REQUIRED_H2 = [
    "标识", "系统概述", "文档概述", "要求的状态和方式", "CSCI能力需求",
    "CSCI外部接口需求", "CSCI内部接口需求", "CSCI内部数据需求", "适应性需求",
    "保密性需求", "安全性需求", "CSCI环境适应性需求", "其他质量特性",
    "计算机资源需求", "设计和实现约束", "人员相关需求", "训练相关需求",
    "软件保障需求", "包装需求", "其他需求", "需求的优先顺序和关键性",
]
SRS_RESOURCE_H3 = ["计算机硬件需求", "计算机硬件资源使用需求", "计算机软件需求", "计算机通信需求"]
SRS_RECORD_HEADERS = ["需求编号", "名称", "适用条件", "需求陈述", "输入", "预期输出", "性能或边界", "异常与恢复", "优先级", "关键性", "合格性方法", "来源编号", "计划证据"]

REQUIRED_H1 = [
    "概述",
    "使用与技术指标要求",
    "总体架构",
    "分系统设计",
    "系统集成设计",
    "关键技术分析",
    "效能分析",
    "工程组织管理",
    "初步工作计划",
]

OVERALL_DOCUMENT_TYPE = "overall-technical-solution"
GJB_DOCUMENT_TYPES = {
    "SDP": ("软件开发计划", "A"),
    "SIP": ("软件安装计划", "B"),
    "STrP": ("软件移交计划", "C"),
    "STP": ("软件测试计划", "D"),
    "OCD": ("运行方案说明", "E"),
    "SSS": ("系统/子系统规格说明", "F"),
    "IRS": ("接口需求规格说明", "G"),
    "SSDD": ("系统/子系统设计说明", "H"),
    "IDD": ("接口设计说明", "I"),
    "SRS": ("软件需求规格说明", "J"),
    "SDD": ("软件设计说明", "K"),
    "DBDD": ("数据库设计说明", "L"),
    "STD": ("软件测试说明", "M"),
    "STR": ("软件测试报告", "N"),
    "SPS": ("软件产品规格说明", "O"),
    "SVD": ("软件版本说明", "P"),
    "SUM": ("软件用户手册", "Q"),
    "CPM": ("计算机编程手册", "R"),
    "FSM": ("固件保障手册", "S"),
    "SDSR": ("软件研制总结报告", "T"),
}

DEMO_H2 = {
    "概述": ["任务依据", "编制目的", "指导原则", "建设目标", "主要工作", "与其他项目关系", "名词术语", "缩略语"],
    "使用与技术指标要求": ["使用要求", "技术指标要求"],
    "总体架构": ["体系架构设计", "逻辑架构设计", "系统架构设计", "功能架构设计", "数据架构设计", "部署架构设计", "数据交换设计", "关键性能指标设计", "组织运用模式", "技术体制"],
    "分系统设计": ["一体化业务分系统"],
    "系统集成设计": ["集成联试"],
    "关键技术分析": ["关键技术一"],
    "效能分析": ["综合效能分析"],
    "工程组织管理": ["组织机构", "考核验收", "安全保密"],
    "初步工作计划": ["初步工作计划"],
}


def strip_heading_number(title: str) -> str:
    """Headings are auto-numbered by the template's multilevel list."""
    return HEADING_NUMBER_RE.sub("", title).strip()


def normalize_document_type(value: str | None) -> str:
    if not value:
        return OVERALL_DOCUMENT_TYPE
    cleaned = value.strip()
    if cleaned.lower() in {OVERALL_DOCUMENT_TYPE, "overall", "总体技术方案"}:
        return OVERALL_DOCUMENT_TYPE
    for code in GJB_DOCUMENT_TYPES:
        if cleaned.lower() == code.lower():
            return code
    supported = ", ".join([OVERALL_DOCUMENT_TYPE, *GJB_DOCUMENT_TYPES])
    raise ValueError(f"Unsupported document_type {value!r}. Supported: {supported}")


def style_names(document: Document) -> set[str]:
    return {style.name for style in document.styles}


def set_style_font(style: Any, east_asia: str, size_pt: float, bold: bool = False, latin: str = "Times New Roman") -> None:
    style.font.name = latin
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def normalize_srs_styles(document: Document) -> None:
    """Convert P-09's direct-format intent into stable semantic styles."""
    normal = document.styles["Normal"]
    set_style_font(normal, "仿宋", 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Twips(480)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Twips(579)

    heading_specs = {
        "Heading 1": ("黑体", 14),
        "Heading 2": ("楷体", 14),
        "Heading 3": ("仿宋", 12),
        "Heading 4": ("仿宋", 12),
    }
    for name, (font_name, size) in heading_specs.items():
        style = document.styles[name]
        set_style_font(style, font_name, size, bold=True)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.first_line_indent = Twips(0)
        style.element.get_or_add_pPr().find(qn("w:ind")).set(qn("w:firstLineChars"), "0")
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(24)

    def ensure(name: str, base: str, font: str, size: float, bold: bool, align: WD_ALIGN_PARAGRAPH):
        style = document.styles[name] if name in style_names(document) else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[base]
        set_style_font(style, font, size, bold=bold)
        style.paragraph_format.alignment = align
        style.paragraph_format.first_line_indent = Twips(0)
        # Character-based indentation takes precedence over twips and is
        # otherwise inherited from Normal even when firstLine is set to zero.
        indent = style.element.get_or_add_pPr().find(qn("w:ind"))
        for attr in ("leftChars", "rightChars", "hangingChars"):
            indent.attrib.pop(qn("w:" + attr), None)
        indent.set(qn("w:firstLineChars"), "0")
        style.paragraph_format.keep_together = True
        return style

    body_style = ensure("SRS正文", "Normal", "仿宋", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY)
    body_style.paragraph_format.first_line_indent = Twips(480)
    body_style.element.get_or_add_pPr().find(qn("w:ind")).set(qn("w:firstLineChars"), "200")
    ensure("SRS图样式", "SRS正文", "仿宋", 12, False, WD_ALIGN_PARAGRAPH.CENTER).paragraph_format.keep_with_next = True
    ensure("SRS题注", "Normal", "黑体", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER).paragraph_format.keep_with_next = True
    ensure("SRS表头", "Normal", "黑体", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER)
    ensure("SRS表正文", "Normal", "宋体", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER)
    toc_heading = ensure("TOC Heading", "Normal", "黑体", 16, False, WD_ALIGN_PARAGRAPH.CENTER)
    toc_heading.paragraph_format.keep_with_next = True
    toc_heading.paragraph_format.space_before = Pt(0)
    toc_heading.paragraph_format.space_after = Pt(6)
    toc_heading.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    toc_heading.paragraph_format.line_spacing = Pt(24)
    # Word/LibreOffice rebuild this title from its style when updating the TOC.
    # The inherited template style contains a blue accent and left alignment.
    color = toc_heading.element.get_or_add_rPr().find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        toc_heading.element.get_or_add_rPr().append(color)
    color.attrib.clear()
    color.set(qn("w:val"), "000000")
    for name in ("SRS题注", "SRS表头", "SRS表正文"):
        fmt = document.styles[name].paragraph_format
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(18)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)


def require_styles(document: Document, document_type: str) -> None:
    if document_type == "SRS":
        required = {"Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "toc 1", "toc 2", "toc 3"}
    else:
        required = {
            "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5",
            "145正文", "145图样式", "145表头", "145表正文", "Caption",
            "311-目录标题", "toc 1", "toc 2", "toc 3", "编号密级", "文件名称", "单位名称",
        }
    missing = sorted(required - style_names(document))
    if missing:
        raise ValueError(f"Template is missing required styles: {', '.join(missing)}")


def style_id_to_name_map(document: Document) -> dict[str, str]:
    return {style.style_id: style.name for style in document.styles if style.style_id}


def paragraph_style_name(p_element: Any, id_map: dict[str, str]) -> str:
    """Resolve the style name of a raw <w:p> element via its styleId."""
    ppr = p_element.find(qn("w:pPr"))
    if ppr is None:
        return "Normal"
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        return "Normal"
    return id_map.get(pstyle.get(qn("w:val")), "Normal")


def clear_body_content(document: Document) -> None:
    """Remove body content from the first Heading 1 paragraph onward.

    Keeps the cover table, TOC title, TOC field, and the front-matter section
    break. The trailing body-level sectPr is always preserved.
    """
    body = document._body._element
    children = list(body)
    id_map = style_id_to_name_map(document)
    first_heading_index = None
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        if paragraph_style_name(child, id_map) == "Heading 1":
            first_heading_index = index
            break

    if first_heading_index is None:
        # Fallback: template without body headings; clear everything but sectPr.
        first_heading_index = 0

    for child in children[first_heading_index:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def resolve_heading_numbering(document: Document, style_name: str) -> tuple[str, str]:
    """Find a real heading list, respecting inherited style bindings first."""
    level = str(HEADING_LEVEL_BY_STYLE[style_name] - 1)
    root = document.part.numbering_part.element
    abstracts = {node.get(qn("w:abstractNumId")): node for node in root.findall(qn("w:abstractNum"))}
    effective: dict[str, dict[str, Any]] = {}
    for node in root.findall(qn("w:num")):
        num_id = node.get(qn("w:numId"))
        abstract_id = node.find(qn("w:abstractNumId"))
        abstract = abstracts.get(abstract_id.get(qn("w:val"))) if abstract_id is not None else None
        if num_id == "0" or abstract is None:
            continue
        levels = {row.get(qn("w:ilvl")): row for row in abstract.findall(qn("w:lvl"))}
        for override in node.findall(qn("w:lvlOverride")):
            replacement = override.find(qn("w:lvl"))
            if replacement is not None:
                levels[override.get(qn("w:ilvl"))] = replacement
        effective[num_id] = levels
    style = document.styles[style_name]
    visited: set[str] = set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        binding = style.element.find("./" + qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId"))
        if binding is not None:
            num_id = binding.get(qn("w:val"))
            if level not in effective.get(num_id, {}):
                raise ValueError(f"Heading style {style_name} binds missing/disabled numbering {num_id} at level {level}")
            return num_id, level
        style = style.base_style
    linked = []
    candidates = []
    for num_id, levels in effective.items():
        if level not in levels:
            continue
        definition = levels[level]
        linked_style = definition.find(qn("w:pStyle"))
        if linked_style is not None and linked_style.get(qn("w:val")) == document.styles[style_name].style_id:
            linked.append(num_id)
        decimal_hierarchy = True
        for index in range(int(level) + 1):
            row = levels.get(str(index))
            fmt = row.find(qn("w:numFmt")) if row is not None else None
            pattern = row.find(qn("w:lvlText")) if row is not None else None
            tokens = re.findall(r"%([1-9])", pattern.get(qn("w:val"), "")) if pattern is not None else []
            if fmt is None or fmt.get(qn("w:val")) != "decimal" or tokens != [str(n) for n in range(1, index + 2)]:
                decimal_hierarchy = False
                break
        if decimal_hierarchy:
            candidates.append(num_id)
    selected = linked or candidates
    if len(selected) != 1:
        raise ValueError(f"Cannot resolve an unambiguous multilevel list for {style_name}; bind the heading style to a valid list")
    return selected[0], level


def add_styled_paragraph(document: Document, text: str, style_name: str) -> None:
    paragraph = document.add_paragraph(style=style_name)
    if style_name in HEADING_LEVEL_BY_STYLE:
        heading_num_id, heading_level = resolve_heading_numbering(document, style_name)
        ppr = paragraph._p.get_or_add_pPr()
        num_pr = ppr.find(qn("w:numPr"))
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            ppr.append(num_pr)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), heading_level)
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), heading_num_id)
        num_pr.append(ilvl)
        num_pr.append(num_id)
    paragraph.add_run(text)


def set_cell_text(cell: Any, text: Any, style_name: str) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = style_name
    paragraph.add_run("" if text is None else str(text))


def replace_paragraph_text(paragraph: Any, text: str) -> None:
    """Replace visible text while keeping the paragraph style and run format."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_first_nonempty_cell_paragraphs(cell: Any, texts: list[str]) -> None:
    """Replace the texts of the first N non-empty paragraphs in a cell."""
    remaining = list(texts)
    for paragraph in cell.paragraphs:
        if not remaining:
            break
        if paragraph.text.strip():
            replace_paragraph_text(paragraph, remaining.pop(0))


def update_cover_table(document: Document, metadata: dict[str, Any]) -> None:
    """Fill the template cover table in place, preserving cover styles."""
    if not document.tables:
        return
    cover = document.tables[0]

    number = str(metadata.get("number", "")).strip()
    identifier = str(metadata.get("identifier", "")).strip()
    title = str(metadata.get("title", "")).strip()
    unit = str(metadata.get("unit", "")).strip()
    date = str(metadata.get("date", "")).strip()

    for row in cover.rows:
        seen: set[int] = set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cell_text = cell.text.strip()
            if number and "编号" in cell_text and len(cell_text) < 30:
                replace_first_nonempty_cell_paragraphs(cell, [f"编号：{number}"])
            elif identifier and "SJZT" in cell_text:
                replace_first_nonempty_cell_paragraphs(cell, [identifier])
            elif title and "总体技术方案" in cell_text:
                replace_first_nonempty_cell_paragraphs(cell, [title])
            elif "技术总师组" in cell_text and (unit or date):
                texts = [t for t in [unit, date] if t]
                replace_first_nonempty_cell_paragraphs(cell, texts)


def update_srs_cover(document: Document, metadata: dict[str, Any]) -> None:
    """Fill the paragraph-based P-09 cover without inventing approvals."""
    values = {
        0: f"密    级：{metadata.get('classification', '未填写')}",
        1: f"版    本：{metadata.get('version', '未填写')}",
        2: f"阶段标注：{metadata.get('phase', '未填写')}",
        6: str(metadata.get("title", "软件需求规格说明（SRS）")),
        15: f"编    制：{metadata.get('author', '未填写')}",
        16: f"审    核：{metadata.get('reviewer', '未填写')}",
        17: f"标    审：{metadata.get('standards_reviewer', '未填写')}",
        18: f"批    准：{metadata.get('approver', '未填写')}",
        21: str(metadata.get("unit", "未填写")),
        22: str(metadata.get("date", "未填写")),
    }
    for index, value in values.items():
        if index < len(document.paragraphs):
            replace_paragraph_text(document.paragraphs[index], value)


def normalize_srs_headers_and_tables(document: Document, metadata: dict[str, Any]) -> None:
    for section in document.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            for paragraph in header.paragraphs:
                if "软件需求规格说明" in paragraph.text or "计划阶段-06" in paragraph.text:
                    replace_paragraph_text(paragraph, (str(metadata["phase"]) + "-" if metadata.get("phase") else "") + "P-09-软件需求规格说明")
    for table in document.tables:
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for row in table.rows[1:]:
            prevent_row_split(row)
        if table.rows and [cell.text.strip() for cell in table.rows[0].cells] == ["版本", "修改原因", "修改内容", "修改人", "修改日期"] and len(table.rows) > 1:
            values = [
                metadata.get("version", "未填写"),
                metadata.get("change_reason", "未提供修改原因"),
                metadata.get("change_content", "未提供修改内容"),
                metadata.get("author", "未填写"),
                metadata.get("date", "未填写"),
            ]
            for cell, value in zip(table.rows[1].cells, values):
                replace_paragraph_text(cell.paragraphs[0], str(value))
            # Empty stationery rows are useful in a blank template but must not
            # create empty continuation pages in a generated review document.
            for row in list(table.rows)[2:]:
                if not any(cell.text.strip() for cell in row.cells):
                    table._tbl.remove(row._tr)
            apply_table_format(table, "SRS")
            set_table_geometry(table)
            for i, row in enumerate(table.rows):
                for height in row._tr.xpath("./w:trPr/w:trHeight"):
                    height.getparent().remove(height)
                for j, cell in enumerate(row.cells):
                    value = cell.text
                    set_cell_text(cell, value, "SRS表头" if i == 0 else "SRS表正文")
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.first_line_indent = Twips(0)
                        paragraph.paragraph_format.left_indent = Twips(0)
                        paragraph.paragraph_format.right_indent = Twips(0)
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if i and j in (1, 2) else WD_ALIGN_PARAGRAPH.CENTER
                        paragraph.paragraph_format.line_spacing = Pt(18)


def enable_update_fields(document: Document) -> None:
    """Ask Word/WPS to refresh the TOC field and page numbers on open."""
    settings = document.settings.element
    element = settings.find(qn("w:updateFields"))
    if element is None:
        element = settings.makeelement(qn("w:updateFields"))
        settings.append(element)
    element.set(qn("w:val"), "true")


def reset_toc_field(document: Document) -> None:
    """Remove stale TOC cache and retain one updateable Word TOC field."""
    for sdt in document._body._element.iter(qn("w:sdt")):
        instructions = " ".join(node.text or "" for node in sdt.iter(qn("w:instrText")))
        if "TOC" not in instructions.upper():
            continue
        content = sdt.find(qn("w:sdtContent"))
        if content is None:
            continue
        for child in list(content):
            content.remove(child)
        title_paragraph = OxmlElement("w:p")
        title_ppr = OxmlElement("w:pPr")
        title_style = OxmlElement("w:pStyle")
        title_style.set(qn("w:val"), document.styles["TOC Heading"].style_id)
        title_ppr.append(title_style)
        title_paragraph.append(title_ppr)
        title_run = OxmlElement("w:r")
        title_text = OxmlElement("w:t")
        title_text.text = "目  录"
        title_run.append(title_text)
        title_paragraph.append(title_run)
        content.append(title_paragraph)
        paragraph = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), document.styles["toc 1"].style_id)
        ppr.append(pstyle)
        paragraph.append(ppr)
        for kind, value in [
            ("fldChar", "begin"),
            ("instrText", ' TOC \\o "1-3" \\h \\z \\u '),
            ("fldChar", "separate"),
            ("text", "打开文档后请更新目录域"),
            ("fldChar", "end"),
        ]:
            run = OxmlElement("w:r")
            node = OxmlElement(f"w:{'t' if kind == 'text' else kind}")
            if kind == "fldChar":
                node.set(qn("w:fldCharType"), value)
            else:
                node.text = value
                if kind == "instrText":
                    node.set(qn("xml:space"), "preserve")
            run.append(node)
            paragraph.append(run)
        content.append(paragraph)

    # The older overall-solution template stores its TOC directly in body
    # paragraphs. Its outer field ends in a separate paragraph carrying sectPr.
    # Clear only the field/cache and retain paragraph/section properties.
    paragraphs = list(document.paragraphs)
    for start, first in enumerate(paragraphs):
        instructions = " ".join(node.text or "" for node in first._p.iter(qn("w:instrText")))
        if not re.search(r"\bTOC\s", instructions, re.I):
            continue
        depth = 0
        end = None
        for index in range(start, len(paragraphs)):
            for node in paragraphs[index]._p.iter(qn("w:fldChar")):
                kind = node.get(qn("w:fldCharType"))
                if kind == "begin": depth += 1
                elif kind == "end": depth -= 1
            if depth == 0:
                end = index
                break
        if end is None:
            raise ValueError("Template contains an unbalanced direct TOC field")
        replaced_start_ids = {node.get(qn("w:id")) for p in paragraphs[start:end + 1] for node in p._p.iter(qn("w:bookmarkStart"))}
        surviving_start_ids = {node.get(qn("w:id")) for node in document._element.iter(qn("w:bookmarkStart"))} - replaced_start_ids
        boundary_ends = [node for p in paragraphs[start:end + 1] for node in p._p.iter(qn("w:bookmarkEnd")) if node.get(qn("w:id")) in surviving_start_ids]
        for index in range(start, end + 1):
            paragraph = paragraphs[index]
            if index == start or paragraph._p.xpath("./w:pPr/w:sectPr"):
                paragraph.clear()
            else:
                paragraph._p.getparent().remove(paragraph._p)
        append_word_field(first, ' TOC \\o "1-3" \\h \\z \\u ', "打开文档后请更新目录域")
        for node in boundary_ends:
            first._p.append(node)
        break


def append_word_field(paragraph: Any, instruction: str, placeholder: str = "0") -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instr_run._r.append(instr)
    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)
    paragraph.add_run(placeholder)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_caption_paragraph(document: Document, caption: str, style_name: str) -> None:
    match = re.match(r"^(图|表)\s*(?:\d+(?:-\d+)?)?\s*(.*)$", caption.strip())
    if not match:
        add_styled_paragraph(document, caption, style_name)
        return
    label, title = match.groups()
    paragraph = document.add_paragraph(style=style_name)
    paragraph.add_run(f"{label} ")
    chapter = sum(1 for p in document.paragraphs if p.style and p.style.name == "Heading 1")
    # STYLEREF \n returns the heading number rather than its text. Distinct
    # sequence identifiers give each chapter its own counter across renderers.
    append_word_field(paragraph, ' STYLEREF "Heading 1" \\n ', str(chapter))
    paragraph.add_run("-")
    sequence = f"{'Table' if label == '表' else 'Figure'}_{chapter}"
    prior = sum(1 for node in document._element.iter(qn("w:instrText")) if re.search(r"\bSEQ\s+" + re.escape(sequence) + r"\b", node.text or ""))
    append_word_field(paragraph, f" SEQ {sequence} \\* ARABIC ", str(prior + 1))
    if title:
        paragraph.add_run(f" {title}")


def require_page_number_structure(path: Path, document_type: str) -> None:
    """Fail fast if generated output loses template page-number structure."""
    result = AuditResult()
    audit_page_numbering(path, result, document_type=document_type)
    if result.errors:
        raise ValueError("Generated DOCX failed page-number structure audit: " + "; ".join(result.errors))


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def apply_table_format(table: Any, document_type: str) -> None:
    """Apply visible borders, fixed layout, margins, and usable widths."""
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "000000")
    cell_margin = tbl_pr.find(qn("w:tblCellMar"))
    if cell_margin is None:
        cell_margin = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_margin)
    for edge, width in [("top", "60"), ("bottom", "60"), ("left", "80"), ("right", "80")]:
        node = cell_margin.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            cell_margin.append(node)
        node.set(qn("w:w"), width)
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table: Any, total_width: int = 8306) -> None:
    """Synchronize tblW, tblGrid, and every tcW after all rows exist."""
    column_count = len(table.columns)
    if column_count <= 0:
        return
    if column_count == 2:
        widths = [1800, total_width - 1800]
    else:
        base = total_width // column_count
        widths = [base] * column_count
        widths[-1] += total_width - sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def add_regular_table(document: Document, headers: list[str], rows: list[Any], document_type: str) -> None:
    table_header_style = "SRS表头" if document_type == "SRS" else TABLE_HEADER_STYLE
    table_body_style = "SRS表正文" if document_type == "SRS" else TABLE_BODY_STYLE
    table = document.add_table(rows=1, cols=len(headers))
    try:
        table.style = GENERATED_TABLE_STYLE
    except KeyError:
        pass
    table.autofit = False
    apply_table_format(table, document_type)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, table_header_style)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        values = list(row_values)
        for idx in range(len(headers)):
            set_cell_text(row.cells[idx], values[idx] if idx < len(values) else "", table_body_style)
            if document_type == "SRS" and (headers == ["字段", "内容"] and idx == 1 or headers[idx] in {"验证条件", "通过准则", "修改原因", "修改内容"}):
                for paragraph in row.cells[idx].paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_table_geometry(table)


def add_table(document: Document, table_data: dict[str, Any], document_type: str) -> None:
    caption_style = "SRS题注" if document_type == "SRS" else "Caption"
    caption = table_data.get("caption")
    headers = [str(value) for value in table_data.get("headers", [])]
    rows = table_data.get("rows", [])
    if not headers:
        return
    if document_type == "SRS" and (len(headers) > 6 or table_data.get("layout") == "record"):
        for row_index, row_values in enumerate(rows, start=1):
            values = list(row_values)
            suffix = values[0] if values else str(row_index)
            if caption:
                add_caption_paragraph(document, f"{caption}—{suffix}" if len(rows) > 1 else str(caption), caption_style)
            add_regular_table(document, ["字段", "内容"], [[header, values[i] if i < len(values) else ""] for i, header in enumerate(headers)], document_type)
        return
    if caption:
        add_caption_paragraph(document, str(caption), caption_style)
    add_regular_table(document, headers, rows, document_type)


def add_figure(document: Document, figure_data: dict[str, Any], document_type: str) -> None:
    image_path = figure_data.get("path")
    if image_path:
        paragraph = document.add_paragraph(style="SRS图样式" if document_type == "SRS" else FIGURE_STYLE)
        section = document.sections[-1]
        paragraph.add_run().add_picture(str(image_path), width=section.page_width - section.left_margin - section.right_margin)
    else:
        placeholder = str(figure_data.get("placeholder", "")).strip()
        add_styled_paragraph(document, placeholder, "SRS图样式" if document_type == "SRS" else FIGURE_STYLE)
    caption = figure_data.get("caption")
    if caption:
        add_caption_paragraph(document, str(caption), "SRS题注" if document_type == "SRS" else "Caption")


def normalize_sections(content: dict[str, Any]) -> list[dict[str, Any]]:
    sections = content.get("sections", [])
    if not isinstance(sections, list):
        raise ValueError("content JSON field 'sections' must be a list")
    return sections


def write_sections(document: Document, sections: list[dict[str, Any]], document_type: str) -> None:
    body_style = "SRS正文" if document_type == "SRS" else "145正文"
    for section in sections:
        level = int(section.get("level", 1))
        title = strip_heading_number(str(section.get("title", "")))
        if not title:
            continue
        style_name = HEADING_STYLE_BY_LEVEL.get(level, "Heading 5")
        add_styled_paragraph(document, title, style_name)

        for paragraph in section.get("paragraphs", []) or []:
            if isinstance(paragraph, dict):
                text = str(paragraph.get("text", ""))
                role_style = str(paragraph.get("style", body_style))
                add_styled_paragraph(document, text, role_style)
            else:
                add_styled_paragraph(document, str(paragraph), body_style)

        for figure_data in section.get("figures", []) or []:
            add_figure(document, figure_data, document_type)

        for table_data in section.get("tables", []) or []:
            add_table(document, table_data, document_type)


def overall_demo_content() -> dict[str, Any]:
    sections: list[dict[str, Any]] = []
    for chapter_index, h1 in enumerate(REQUIRED_H1, start=1):
        sections.append(
            {
                "level": 1,
                "title": h1,
                "paragraphs": ["本章为 GJB438C 模板格式样例内容，实际编写时应替换为项目真实资料。"],
            }
        )
        for h2 in DEMO_H2[h1]:
            section: dict[str, Any] = {
                "level": 2,
                "title": h2,
                "paragraphs": ["本节作为二级标题格式控制单元，正文统一采用模板中的 145正文 样式。"],
            }
            if h2 == "名词术语":
                section["tables"] = [
                    {
                        "caption": f"表{chapter_index}-1 名词术语",
                        "headers": ["编号", "名称", "说明"],
                        "rows": [["1", "示例术语", "用于验证表格样式。"]],
                    }
                ]
            elif h2 == "缩略语":
                section["tables"] = [
                    {
                        "caption": f"表{chapter_index}-2 缩略语",
                        "headers": ["序号", "简写", "全称", "解释说明"],
                        "rows": [["1", "GJB", "国家军用标准", "用于验证缩略语表格式。"]],
                    }
                ]
            elif h2 == "体系架构设计":
                section["figures"] = [{"caption": f"图 {chapter_index}-1 体系架构示意图"}]
            sections.append(section)
    return {"metadata": {"document_type": OVERALL_DOCUMENT_TYPE, "title": "总体技术方案"}, "sections": sections}


def srs_demo_content() -> dict[str, Any]:
    """Create a complete Appendix-J skeleton with reviewable record examples."""
    req_headers = ["需求编号", "名称", "适用条件", "需求陈述", "输入", "预期输出", "性能或边界", "异常与恢复", "优先级", "关键性", "合格性方法", "来源编号", "计划证据"]

    def req_row(rid: str, name: str, condition: str, statement: str, priority: str, criticality: str, method: str, source: str) -> list[str]:
        return [rid, name, condition, statement, "受控请求或配置", "可判定的业务结果或状态", "边界由批准的上层需求规定", "无效输入或依赖异常时返回受控错误并记录诊断信息", priority, criticality, method, source, f"{rid}-验证记录"]
    sections: list[dict[str, Any]] = [
        {"level": 1, "title": "范围", "paragraphs": ["本章规定本文档的适用对象、系统概况和文档使用范围。"]},
        {"level": 2, "title": "标识", "paragraphs": ["CSCI名称、标识号、版本和发布号均应以批准的配置项清单为准；未批准项列入TBD登记。"]},
        {"level": 2, "title": "系统概述", "paragraphs": ["本节说明CSCI用途、规模、关键特性、相关方、运行现场和生命周期背景。"]},
        {"level": 2, "title": "文档概述", "paragraphs": ["本文档规定CSCI验收需求、合格性方法和需求可追踪关系。"]},
        {"level": 1, "title": "引用文档", "paragraphs": ["引用文件的版本和适用性由项目配置管理控制。"], "tables": [{
            "caption": "表2-1 引用文档清单", "headers": ["文档编号", "标题", "编写单位", "修订版", "日期", "获取来源"],
            "rows": [["GJB 438C-2021", "军用软件开发文档通用要求", "标准发布机构", "2021", "2021", "正常渠道"]]
        }]},
        {"level": 1, "title": "需求", "paragraphs": ["本章中的正式需求均具有项目唯一标识、合格性方法和来源或派生依据。"]},
        {"level": 2, "title": "要求的状态和方式", "paragraphs": ["本示例仅定义正常运行方式；真实项目应列出会改变需求的全部状态和方式。"]},
        {"level": 2, "title": "CSCI能力需求", "paragraphs": ["能力按项目批准的功能分解组织。"]},
        {"level": 3, "title": "示例能力", "paragraphs": ["本条展示正式需求记录的最小字段。"], "tables": [{
            "caption": "表3-1 示例能力需求", "headers": req_headers,
            "rows": [req_row("SRS-CAP-DEMO-001", "记录查询", "正常运行方式且用户已通过鉴权", "系统应按授权范围返回满足查询条件的记录，并在输入无效时返回可识别的错误信息。", "高", "一般", "演示、测试", "SYS-REQ-001")]
        }]},
        {"level": 2, "title": "CSCI外部接口需求", "paragraphs": ["外部系统行为仅作为前置条件或触发条件描述。"]},
        {"level": 3, "title": "接口标识和接口图", "paragraphs": ["正式文档应给出接口上下文图；本结构样例以受控接口清单表示。"], "tables": [{
            "caption": "表3-2 外部接口清单", "headers": ["接口编号", "CSCI端", "外部端", "方向", "输入输出/数据", "协议", "鉴权", "异常处理", "前置条件", "验证证据", "状态"],
            "rows": [["SRS-IF-EXT-001", "示例CSCI", "外部身份源", "双向", "身份校验请求/响应及错误码", "HTTPS/JSON", "双向身份校验", "超时、无效响应和鉴权失败返回受控错误", "接口规范已批准且外部服务可用", "接口测试记录和审查记录", "待项目确认"]]
        }]},
        {"level": 3, "title": "外部身份源接口", "paragraphs": ["前置条件为外部身份源按已批准协议提供服务；CSCI应校验请求和响应，记录审计信息，并在超时或鉴权失败时返回受控错误。"], "tables": [{
            "caption": "表3-3 外部接口需求", "headers": req_headers,
            "rows": [req_row("SRS-IF-EXT-001", "身份源接入", "接口规范已批准且外部服务可用", "CSCI应通过批准的接口完成身份校验，并对超时、无效响应和鉴权失败进行受控处理。", "高", "保密关键", "测试、审查", "SYS-IF-001")]
        }]},
    ]
    simple_h2 = [
        ("CSCI内部接口需求", "SRS-IF-INT-001", "内部组件调用", "CSCI内部组件应按受控契约交换数据并记录异常。", "审查、测试", "SYS-ARCH-001"),
        ("CSCI内部数据需求", "SRS-DAT-001", "数据完整性", "CSCI应校验受控数据对象的必填字段和完整性约束。", "测试", "SYS-DAT-001"),
        ("适应性需求", "SRS-ADP-001", "参数适配", "CSCI应通过受控配置适配部署环境参数并记录配置变更审计信息。", "演示、审查", "SYS-ADP-001"),
        ("保密性需求", "SRS-SEC-001", "访问控制", "CSCI应依据用户身份和授权策略限制受保护资源访问并记录安全审计。", "测试、审查", "SYS-SEC-001"),
        ("安全性需求", "SRS-SAF-001", "危险操作保护", "CSCI应在高后果操作执行前校验前置条件并提供受控确认。", "演示、测试", "SYS-SAF-001"),
        ("CSCI环境适应性需求", "SRS-ENV-001", "运行环境", "CSCI应在批准的硬件和操作系统基线中完成安装、启动和基本运行。", "测试、审查", "SYS-ENV-001"),
        ("其他质量特性", "SRS-QUA-001", "可恢复性", "CSCI应在批准的恢复点和恢复时间目标内恢复受保护服务。", "测试、分析", "SYS-QUA-001"),
    ]
    for title, rid, name, statement, method, source in simple_h2:
        sections.append({"level": 2, "title": title, "tables": [{"caption": f"表 {title}示例", "headers": req_headers, "rows": [req_row(rid, name, "适用条件由项目基线定义", statement, "高", "一般", method, source)]}]})
    sections.append({"level": 2, "title": "计算机资源需求", "paragraphs": ["资源需求给出可测量条件和批准基线。"]})
    resource_rows = [
        ("计算机硬件需求", "SRS-RES-HW-001", "硬件基线", "CSCI应能够在批准的计算机硬件基线上部署和运行。", "审查、测试", "SYS-RES-001"),
        ("计算机硬件资源使用需求", "SRS-RES-USE-001", "资源利用", "CSCI应在规定负载条件下将资源使用保持在批准阈值内。", "测试、分析", "SYS-RES-002"),
        ("计算机软件需求", "SRS-RES-SW-001", "基础软件", "CSCI应与批准的软件运行环境和版本基线兼容。", "审查、测试", "SYS-RES-003"),
        ("计算机通信需求", "SRS-RES-COM-001", "通信资源", "CSCI应在批准的网络拓扑、带宽和可用时间条件下交换业务数据。", "测试、分析", "SYS-RES-004"),
    ]
    for title, rid, name, statement, method, source in resource_rows:
        sections.append({"level": 3, "title": title, "tables": [{"caption": f"表 {title}示例", "headers": req_headers, "rows": [req_row(rid, name, "批准的运行环境", statement, "高", "一般", method, source)]}]})
    trailing = [
        ("设计和实现约束", "SRS-CST-001", "设计约束", "CSCI应遵循批准的体系结构、编码和数据标准。", "审查", "SYS-CST-001"),
        ("人员相关需求", "SRS-PER-001", "在线用户", "CSCI应在规定并发用户条件下保持已定义的服务能力。", "测试、分析", "SYS-PER-001"),
        ("训练相关需求", "SRS-TRN-001", "内置帮助", "CSCI应向授权用户提供与其角色相符的在线帮助信息。", "演示、审查", "SYS-TRN-001"),
        ("软件保障需求", "SRS-SUP-001", "维护诊断", "CSCI应提供用于故障定位的受控日志和诊断信息。", "演示、审查", "SYS-SUP-001"),
        ("包装需求", "SRS-PKG-001", "交付包装", "CSCI应以批准的介质、目录结构和标识信息形成交付包。", "审查", "SYS-PKG-001"),
        ("其他需求", "SRS-OTH-001", "其他约束", "CSCI应保存项目批准的其他验收约束及其变更记录。", "审查", "SYS-OTH-001"),
        ("需求的优先顺序和关键性", "SRS-PRI-001", "优先级标识", "CSCI需求记录应标识优先级、关键性以及保密或安全相关属性。", "审查", "SYS-PRI-001"),
    ]
    all_ids = ["SRS-CAP-DEMO-001", "SRS-IF-EXT-001"] + [row[1] for row in simple_h2 + resource_rows + trailing]
    all_sources = ["SYS-REQ-001", "SYS-IF-001"] + [row[5] for row in simple_h2 + resource_rows + trailing]
    all_methods = ["演示、测试", "测试、审查"] + [row[4] for row in simple_h2 + resource_rows + trailing]
    for title, rid, name, statement, method, source in trailing:
        sections.append({"level": 2, "title": title, "tables": [{"caption": f"表 {title}示例", "headers": req_headers, "rows": [req_row(rid, name, "适用条件由项目基线定义", statement, "中", "一般", method, source)]}]})
    sections.extend([
        {"level": 1, "title": "合格性规定", "paragraphs": ["合格性方法包括演示、测试、分析、审查和特殊合格性方法。"], "tables": [{
            "caption": "表4-1 需求合格性矩阵", "headers": ["需求编号", "合格性方法", "验证条件", "计划证据", "通过准则"],
            "rows": [[rid, method, "批准的验证环境", f"验证记录-{i:03d}", "满足需求陈述和边界"] for i, (rid, method) in enumerate(zip(all_ids, all_methods), 1)]
        }]},
        {"level": 1, "title": "需求可追踪性", "paragraphs": ["本章同时给出来源到SRS及SRS到来源的追踪关系。"], "tables": [
            {"caption": "表5-1 来源到SRS正向追踪", "headers": ["来源编号", "来源分类", "SRS需求编号", "落实位置", "处置"], "rows": [[source, "示例来源", rid, "第3章", "落实"] for source, rid in zip(all_sources, all_ids)]},
            {"caption": "表5-2 SRS到来源反向追踪", "headers": ["SRS需求编号", "类别", "来源编号或派生依据", "合格性方法", "计划证据"], "rows": [[rid, "示例类别", source, method, f"验证记录-{i:03d}"] for i, (rid, source, method) in enumerate(zip(all_ids, all_sources, all_methods), 1)]},
        ]},
        {"level": 1, "title": "注释", "paragraphs": ["本章说明缩略语、术语、TBD登记及必要的背景信息。"], "tables": [{
            "caption": "表6-1 TBD登记", "headers": ["TBD编号", "未决事项", "影响范围", "责任方", "关闭条件", "状态"],
            "rows": [["TBD-001", "项目正式标识待确认", "封面与1.1", "项目负责人", "配置项清单批准", "打开"]]
        }]},
    ])
    return {"metadata": {"document_type": "SRS", "title": "软件需求规格说明（SRS）", "classification": "待确认", "version": "V0.1", "phase": "计划阶段", "unit": "待确认", "date": "待确认"}, "sections": sections}


def standard_gjb_demo_content(document_type: str) -> dict[str, Any]:
    name, appendix = GJB_DOCUMENT_TYPES[document_type]
    sections = [
        {
            "level": 1,
            "title": "范围",
            "paragraphs": [
                f"本章说明{name}的适用范围、编制对象和使用边界。该样例用于验证标准 GJB 文档类型不会被强制套用总体技术方案九章结构。"
            ],
        },
        {
            "level": 1,
            "title": "引用文件",
            "paragraphs": ["本章列出编制本文件所依据的标准、合同、任务书和上级文件。"],
            "tables": [
                {
                    "caption": "表2-1 引用文件清单",
                    "headers": ["序号", "文件名称", "说明"],
                    "rows": [["1", "GJB 438C-2021", f"对应附录 {appendix} 的{name}格式。"]],
                }
            ],
        },
        {
            "level": 1,
            "title": f"{name}正文",
            "paragraphs": [
                "本章承载该文档类型的主体内容。真实生成时应按 GJB 438C-2021 对应附录逐条展开。"
            ],
        },
        {
            "level": 2,
            "title": "裁剪说明",
            "paragraphs": ["如项目生命周期或合同约定不适用某些条款，应标注本条无内容并说明理由。"],
        },
        {
            "level": 1,
            "title": "附录",
            "paragraphs": ["本章用于承载可独立维护的图表、清单、数据或补充说明。"],
        },
    ]
    return {
        "metadata": {
            "document_type": document_type,
            "title": name,
        },
        "sections": sections,
    }


def demo_content(document_type: str) -> dict[str, Any]:
    if document_type == OVERALL_DOCUMENT_TYPE:
        return overall_demo_content()
    if document_type == "SRS":
        return srs_demo_content()
    return standard_gjb_demo_content(document_type)


def display_value(value: Any) -> str:
    """Preserve supplied composite structures in the generated document."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2)
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def canonical_requirement_row(requirement: dict[str, Any]) -> list[str]:
    """Legacy helper retained for callers using the old thirteen-column table."""
    mapping = dict(zip(SRS_RECORD_HEADERS, [
        requirement.get("id", ""), requirement.get("name", ""), requirement.get("applicability", ""),
        requirement.get("statement", ""), requirement.get("input", ""), requirement.get("output", ""),
        requirement.get("boundary", ""), requirement.get("exception", ""), requirement.get("priority", ""),
        requirement.get("criticality", ""), "、".join(requirement.get("qualification_methods", [])),
        "、".join(requirement.get("source_ids", [])) or requirement.get("source", ""), requirement.get("evidence", ""),
    ]))
    return [display_value(mapping[header]) for header in SRS_RECORD_HEADERS]


def canonical_srs_to_sections(content: dict[str, Any], mode: str = "draft") -> list[dict[str, Any]]:
    """Render validated source facts; never infer approvals or source coverage."""
    content, _ = require_valid_srs_content(content, mode)
    metadata = content["metadata"]
    requirements = content["requirements"]
    by_clause: dict[str, list[dict[str, Any]]] = {}
    for req in requirements:
        by_clause.setdefault(req["clause"], []).append(req)
    tailoring = {row["clause"]: row for row in content.get("tailoring", [])}
    labels = {
        "id": "标识", "name": "名称", "definition": "定义方式", "entry": "进入条件", "exit": "退出条件",
        "mode": "方式", "description": "说明", "transitions": "状态转换", "purpose": "用途", "parent_id": "上级能力编号",
        "type": "类型", "owner": "责任方", "classification": "分级", "lifecycle": "生命周期", "integrity": "完整性",
        "retention": "保留要求", "backup_recovery": "备份恢复", "migration": "迁移", "fields": "数据元素", "constraints": "约束",
        "csci_endpoint": "CSCI端", "external_endpoint": "外部端", "direction": "方向", "data": "输入输出/数据",
        "protocol": "协议", "authentication": "鉴权", "exception": "异常处理", "precondition": "前置条件", "evidence": "计划证据",
        "status": "状态", "timing": "时序", "capacity": "容量", "version": "版本", "security": "保密性", "kind": "接口类别",
        "irs_reference": "受控IRS引用", "statement": "陈述", "reference_id": "引用文件编号", "location": "定位", "allocation": "分配范围",
        "disposition": "处置", "reason": "原因", "number": "文档编号", "title": "标题", "organization": "编写单位",
        "revision": "修订版", "date": "日期", "source": "获取来源", "normative": "规范性引用", "approval": "受控批准依据",
        "csci_id": "配置项标识", "document_id": "文档编号", "release": "发布号", "issue": "未决事项", "impact": "影响范围", "closure_condition": "关闭条件", "closure_evidence": "关闭证据", "affected_ids": "受影响对象", "due": "期限",
    }

    def record_table(row: dict[str, Any], caption: str, overrides: dict[str, str] | None = None) -> dict[str, Any]:
        names = dict(labels)
        names.update(overrides or {})
        return {"caption": "表 " + caption, "headers": [names.get(key, key) for key in row], "rows": [[display_value(value) for value in row.values()]], "layout": "record"}

    def req_tables(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
        output = []
        req_labels = {"id": "需求编号", "clause": "所属条款", "statement": "需求陈述", "applicability": "适用条件", "input": "输入", "output": "预期输出", "boundary": "性能或边界", "exception": "异常与恢复", "priority": "优先级", "criticality": "关键性", "qualification_methods": "合格性方法", "source_ids": "来源编号", "source": "来源编号", "derivation": "派生依据", "state_ids": "适用状态", "capability_id": "能力编号", "capability": "能力名称", "interface_id": "接口编号", "data_ids": "数据编号"}
        for req in items:
            row = dict(req)
            if row.get("source_ids"):
                row.pop("source", None)
            for key in ("source_ids", "qualification_methods", "state_ids", "data_ids"):
                if key in row:
                    row[key] = "、".join(row[key])
            output.append(record_table(row, f"{req['id']} 需求记录", req_labels))
        return output

    def clause_section(clause: str, *, has_content: bool = False) -> dict[str, Any]:
        row: dict[str, Any] = {"level": 3 if clause.count(".") == 2 else 2, "title": CLAUSE_TITLES[clause], "paragraphs": [], "tables": req_tables(by_clause.get(clause, []))}
        has_content = has_content or bool(row["tables"])
        disposition = tailoring.get(clause, {})
        if clause.startswith("3.11.") and tailoring.get("3.11", {}).get("status") == "not_applicable":
            disposition = tailoring["3.11"]
        if disposition.get("status") == "not_applicable":
            row["paragraphs"].append(("本条无内容。适用性继承 3.11；原因：" if clause.startswith("3.11.") and tailoring.get("3.11", {}).get("status") == "not_applicable" else "本条无内容。原因：") + disposition.get("reason", "未提供（草稿）") + "；裁剪依据：" + disposition.get("basis", "未提供（草稿）"))
        elif not has_content:
            row["paragraphs"].append("尚未提供本条内容或适用性判定（草稿，未作裁剪结论）。")
            if disposition.get("reason"):
                row["paragraphs"].append(disposition["reason"])
        return row

    sections: list[dict[str, Any]] = [
        {"level": 1, "title": "范围"},
        {"level": 2, "title": "标识", "paragraphs": ([metadata["identification"]] if metadata.get("identification") else []) + ["；".join(f"{labels.get(key, key)}：{metadata[key]}" for key in ("csci_id", "version", "release", "document_id", "status") if metadata.get(key)) or "尚未提供配置项标识（草稿）。"]},
        {"level": 2, "title": "系统概述", "paragraphs": [metadata.get("system_overview") or "尚未提供系统用途、规模、相关方及运行现场（草稿）。"]},
        {"level": 2, "title": "文档概述", "paragraphs": [metadata.get("document_overview") or "尚未提供文档用途及使用范围（草稿）。"]},
        {"level": 1, "title": "引用文档", "tables": [record_table(row, row["number"] + " 引用文档") for row in content.get("references", [])]},
        {"level": 1, "title": "需求"},
    ]
    states = content.get("states", [])
    row = clause_section("3.1", has_content=bool(states))
    row["tables"] = [record_table(item, item["id"] + " 状态和方式") for item in states] + row["tables"]
    sections.append(row)

    capabilities = content.get("capabilities", [])
    sections.append(clause_section("3.2", has_content=bool(capabilities)))
    sections[-1]["tables"] = []  # Place each requirement exactly once in its capability group.
    grouped: set[str] = set()
    for cap in capabilities:
        items = [req for req in by_clause.get("3.2", []) if req.get("capability_id") == cap["id"]]
        grouped.update(req["id"] for req in items)
        sections.append({"level": 3, "title": cap["name"] + "（" + cap["id"] + "）", "tables": [record_table(cap, cap["id"] + " 能力定义")] + req_tables(items)})
    legacy_groups: dict[str, list[dict[str, Any]]] = {}
    for req in by_clause.get("3.2", []):
        if req["id"] not in grouped:
            legacy_groups.setdefault(req.get("capability") or "能力需求", []).append(req)
    for name, items in legacy_groups.items():
        sections.append({"level": 3, "title": name, "tables": req_tables(items)})

    interfaces = content.get("interfaces", [])
    for clause, kind in (("3.3", "external"), ("3.4", "internal")):
        selected = [interface for interface in interfaces if interface.get("kind", "external") == kind]
        row = clause_section(clause, has_content=bool(selected))
        row["tables"] = []
        sections.append(row)
        if clause == "3.3" and (selected or by_clause.get(clause)):
            diagram = content.get("interface_diagram")
            context: dict[str, Any] = {"level": 3, "title": "接口标识和接口图", "paragraphs": []}
            if isinstance(diagram, dict) and diagram.get("path"):
                context["figures"] = [{"path": diagram["path"], "caption": "图 接口上下文图"}]
                if diagram.get("description"): context["paragraphs"].append(diagram["description"])
            elif diagram:
                context["paragraphs"].append(display_value(diagram))
            elif any(interface.get("definition") == "irs_reference" for interface in selected):
                context["paragraphs"].append("接口定义引用各接口登记所列受控 IRS。")
            else:
                context["paragraphs"].append("接口标识见各接口登记；接口关系图或其受控引用尚未提供（草稿）。")
            sections.append(context)
        grouped = set()
        for interface in selected:
            items = [req for req in by_clause.get(clause, []) if req.get("interface_id") == interface["id"]]
            grouped.update(req["id"] for req in items)
            sections.append({"level": 3, "title": interface["name"] + "（" + interface["id"] + "）", "tables": [record_table(interface, interface["id"] + " 接口定义", {"id": "接口编号"})] + req_tables(items)})
        remaining = [req for req in by_clause.get(clause, []) if req["id"] not in grouped]
        if remaining:
            sections.append({"level": 3, "title": "未关联接口的需求（草稿）", "tables": req_tables(remaining)})

    data = content.get("data", [])
    row = clause_section("3.5", has_content=bool(data))
    row["tables"] = [record_table(item, item["id"] + " 数据定义") for item in data] + row["tables"]
    sections.append(row)
    for number in range(6, 12):
        clause = f"3.{number}"
        resource_content = number == 11 and any(by_clause.get(f"3.11.{i}") for i in range(1, 5))
        sections.append(clause_section(clause, has_content=resource_content))
        if number == 11:
            for i in range(1, 5): sections.append(clause_section(f"3.11.{i}"))
    individual_priorities = bool(requirements) and all(req.get("priority") and req.get("criticality") for req in requirements) and tailoring.get("3.18", {}).get("status") != "not_applicable"
    for number in range(12, 19):
        sections.append(clause_section(f"3.{number}", has_content=number == 18 and (bool(content.get("priority_policy")) or individual_priorities)))
        if number == 18 and individual_priorities and not content.get("priority_policy"):
            sections[-1]["paragraphs"].append("需求的优先顺序和关键性见各正式需求记录中的对应属性。")
        if number == 18 and content.get("priority_policy"):
            sections[-1]["tables"].insert(0, record_table(content["priority_policy"], "需求优先顺序和关键性策略", {"mode": "分配方式", "statement": "全局说明"}))

    def target_text(item: dict[str, Any]) -> str:
        if item.get("external_target"):
            target = item["external_target"]
            return "IRS引用：" + target["reference_id"] + " / " + target["requirement_id"] + " / " + target["location"]
        return item["requirement_id"]

    qualifications = content.get("qualification", [])
    sections.append({"level": 1, "title": "合格性规定", "tables": [{"caption": "表 需求合格性矩阵", "headers": ["需求编号", "合格性方法", "验证条件", "计划证据", "通过准则"], "rows": [[target_text(item), "、".join(item["methods"]), item.get("condition", ""), item.get("evidence", ""), item.get("pass_criterion", "")] for item in qualifications]}] if qualifications else [], "paragraphs": [] if qualifications else ["尚未提供合格性规定（草稿）。"]})
    sections[-1]["tables"].extend(record_table(item, item["name"] + " 合格性方法扩展") for item in content.get("qualification_method_extensions", []))
    forward, reverse = content.get("forwardTrace", []), content.get("reverseTrace", [])
    trace_tables = [record_table(item, item["id"] + " 来源需求台账", {"id": "来源编号"}) for item in content.get("source_requirements", [])]
    if forward:
        trace_tables.append({"caption": "表 来源到SRS正向追踪", "headers": ["来源编号", "来源分类", "SRS需求编号", "落实位置", "处置"], "rows": [[item["source"], item.get("classification", ""), target_text(item), item.get("location", ""), item.get("disposition", "")] for item in forward]})
    if reverse:
        trace_tables.append({"caption": "表 SRS到来源反向追踪", "headers": ["SRS需求编号", "类别", "来源编号或派生依据", "合格性方法", "计划证据"], "rows": [[target_text(item), item.get("category", ""), "、".join(item.get("source_ids", [])) or item.get("source", "") or item.get("derivation_basis", ""), "、".join(item.get("methods", [])), item.get("evidence", "")] for item in reverse]})
    sections.append({"level": 1, "title": "需求可追踪性", "tables": trace_tables, "paragraphs": [] if trace_tables else ["尚未提供独立来源台账及追踪关系（草稿）。"]})
    sections.append({"level": 1, "title": "注释", "paragraphs": [content["notes"]] if content.get("notes") else [], "tables": [record_table(item, item["id"] + " TBD登记", {"id": "TBD编号"}) for item in content.get("tbd", [])]})
    return sections


def _validate_lowlevel_sections(content: dict[str, Any], mode: str) -> None:
    if mode == "review":
        raise ValueError("Review mode requires the canonical SRS model; low-level sections cannot bypass content validation")
    if content.get("requirements"):
        raise ValueError("Do not combine canonical requirements with low-level sections: supplied requirements would be ignored")
    sections = content.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ValueError("Low-level sections must be a non-empty list")
    for i, section in enumerate(sections):
        if not isinstance(section, dict) or section.get("level", 1) not in range(1, 6) or not isinstance(section.get("title"), str) or not section["title"].strip():
            raise ValueError(f"Invalid low-level section at /sections/{i}")
        for key in ("paragraphs", "tables", "figures"):
            if key in section and not isinstance(section[key], list):
                raise ValueError(f"/sections/{i}/{key} must be an array")
        for j, table in enumerate(section.get("tables", [])):
            if not isinstance(table, dict) or not isinstance(table.get("headers"), list) or not table["headers"] or not isinstance(table.get("rows", []), list):
                raise ValueError(f"Invalid low-level table at /sections/{i}/tables/{j}")
            if any(not isinstance(row, list) or len(row) != len(table["headers"]) for row in table.get("rows", [])):
                raise ValueError(f"Table row width mismatch at /sections/{i}/tables/{j}")


def load_content(path: Path | None, use_demo: bool, document_type_override: str | None, mode: str = "draft") -> dict[str, Any]:
    document_type = normalize_document_type(document_type_override)
    if use_demo:
        content = demo_content(document_type)
    else:
        if not path:
            raise ValueError("Either --content-json or --demo is required")
        content = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(content, dict):
        raise ValueError("Content JSON must be an object")
    metadata = content.setdefault("metadata", {})
    if not isinstance(metadata, dict):
        raise ValueError("Content metadata must be an object")
    if document_type_override:
        declared_type = metadata.get("document_type")
        if declared_type and normalize_document_type(declared_type) != document_type:
            raise ValueError("--document-type conflicts with metadata.document_type")
        metadata["document_type"] = document_type
    else:
        metadata["document_type"] = normalize_document_type(str(metadata.get("document_type", OVERALL_DOCUMENT_TYPE)))
    if metadata["document_type"] == "SRS":
        if content.get("sections"):
            _validate_lowlevel_sections(content, mode)
        else:
            content, _ = require_valid_srs_content(content, mode)
    elif not content.get("sections"):
        raise ValueError("Content produced no sections; refuse to generate an empty document")
    return content


def build_document(template: Path, content: dict[str, Any], output: Path, mode: str = "draft") -> None:
    if mode not in {"draft", "review"}:
        raise ValueError("mode must be draft or review")
    document_type = normalize_document_type(str((content.get("metadata") or {}).get("document_type", OVERALL_DOCUMENT_TYPE)))
    if document_type == "SRS" and content.get("sections"):
        _validate_lowlevel_sections(content, mode)
        sections = normalize_sections(content)
    elif document_type == "SRS":
        content, _ = require_valid_srs_content(content, mode)
        sections = canonical_srs_to_sections(content, mode)
    else:
        sections = normalize_sections(content)
    document = Document(template)
    require_styles(document, document_type)
    if document_type == "SRS":
        normalize_srs_styles(document)
    clear_body_content(document)
    if document_type == "SRS":
        update_srs_cover(document, content.get("metadata", {}) or {})
        normalize_srs_headers_and_tables(document, content.get("metadata", {}) or {})
    else:
        update_cover_table(document, content.get("metadata", {}) or {})
    write_sections(document, sections, document_type)
    enable_update_fields(document)
    reset_toc_field(document)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=".srs-", suffix=".docx", dir=output.parent, delete=False) as stream:
            temporary = Path(stream.name)
        document.save(temporary)
        with zipfile.ZipFile(temporary) as package:
            bad_part = package.testzip()
            if bad_part:
                raise ValueError(f"Generated DOCX has a corrupt package part: {bad_part}")
        Document(temporary)
        require_page_number_structure(temporary, document_type)
        os.replace(temporary, output)
        temporary = None
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--content-json", type=Path, help="Structured JSON content file.")
    source.add_argument("--demo", action="store_true", help="Generate a draft structure demonstration; not review certification.")
    parser.add_argument("--document-type", help="Document type: overall-technical-solution or one of SDP/SIP/STrP/STP/OCD/SSS/IRS/SSDD/IDD/SRS/SDD/DBDD/STD/STR/SPS/SVD/SUM/CPM/FSM/SDSR.")
    parser.add_argument("--template", type=Path, help="Template DOCX path. Defaults by document type.")
    parser.add_argument("--output", type=Path, required=True, help="Output DOCX path.")
    parser.add_argument("--mode", choices=["draft", "review"], default="draft", help="Draft reports missing facts; review blocks unresolved model requirements. Neither mode certifies visual review.")
    parser.add_argument("--json", action="store_true", help="Emit a machine-readable generation result.")
    return parser.parse_args()


def main() -> int:
    # Reports contain Chinese text even when stdout is redirected on Windows.
    # Define an explicit wire encoding instead of inheriting a legacy code page.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    args = parse_args()
    issues: list[dict[str, Any]] = []
    try:
        content = load_content(args.content_json, args.demo, args.document_type, args.mode)
        document_type = normalize_document_type(str((content.get("metadata") or {}).get("document_type", OVERALL_DOCUMENT_TYPE)))
        if document_type == "SRS":
            if content.get("sections"):
                issues.append({"rule_id": "SRS-LOWLEVEL-DRAFT", "severity": "warning", "source": "project-model", "clause": "", "object_id": "", "location": "/sections", "message": "低层草稿装配未经过规范模型内容审核"})
            else:
                issues = validate_srs_content(content, args.mode)
        template = args.template or (SRS_TEMPLATE if document_type == "SRS" else DEFAULT_TEMPLATE)
        build_document(template, content, args.output, args.mode)
    except (OSError, ValueError, TypeError, KeyError, zipfile.BadZipFile) as exc:
        issues = getattr(exc, "issues", issues)
        failure = {"ok": False, "mode": args.mode, "output": None, "issues": issues, "error": str(exc), "checks": {"semantic_review": "not_run", "visual_review": "not_run"}}
        print(json.dumps(failure, ensure_ascii=False, indent=2) if args.json else str(exc), file=sys.stdout if args.json else sys.stderr)
        return 1 if issues and not any(row["rule_id"].startswith("SRS-SCHEMA") for row in issues) else 2
    result = {"ok": True, "mode": args.mode, "output": str(args.output), "issues": issues, "checks": {"model": "not_run" if content.get("sections") else "checked", "package": "passed", "semantic_review": "not_run", "visual_review": "not_run"}}
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        for issue in issues:
            print(f"{issue['severity'].upper()} {issue['rule_id']} {issue['location']}: {issue['message']}", file=sys.stderr)
        print(f"Wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
