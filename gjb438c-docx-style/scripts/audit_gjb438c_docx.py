#!/usr/bin/env python3
"""Profile-aware GJB 438C DOCX and SRS semantic audit."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.oxml.ns import qn


SKILL_DIR = Path(__file__).resolve().parents[1]
def load_document_profiles() -> dict:
    # Loading must not prevent --validate-skill from diagnosing a broken bundle.
    return json.loads((SKILL_DIR / "references" / "document-profiles.json").read_text(encoding="utf-8"))



REQUIRED_STYLES = {
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Heading 4",
    "Heading 5",
    "145正文",
    "145图样式",
    "145表头",
    "145表正文",
    "Caption",
    "311-目录标题",
    "toc 1",
    "toc 2",
    "toc 3",
    "编号密级",
    "文件名称",
    "单位名称",
}

SRS_REQUIRED_STYLES = {"Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "toc 1", "toc 2", "toc 3", "SRS正文", "SRS图样式", "SRS题注", "SRS表头", "SRS表正文"}
SRS_REQUIRED_H1 = ["范围", "引用文档", "需求", "合格性规定", "需求可追踪性", "注释"]
SRS_REQUIRED_H2 = ["标识", "系统概述", "文档概述", "要求的状态和方式", "CSCI能力需求", "CSCI外部接口需求", "CSCI内部接口需求", "CSCI内部数据需求", "适应性需求", "保密性需求", "安全性需求", "CSCI环境适应性需求", "其他质量特性", "计算机资源需求", "设计和实现约束", "人员相关需求", "训练相关需求", "软件保障需求", "包装需求", "其他需求", "需求的优先顺序和关键性"]
SRS_RESOURCE_H3 = ["计算机硬件需求", "计算机硬件资源使用需求", "计算机软件需求", "计算机通信需求"]
SRS_FORMAL_HEADERS = ["需求编号", "名称", "适用条件", "需求陈述", "输入", "预期输出", "性能或边界", "异常与恢复", "优先级", "关键性", "合格性方法", "来源编号", "计划证据"]
QUALIFICATION_METHODS = {"演示", "测试", "分析", "审查", "特殊合格性方法"}
REQ_ID_RE = re.compile(r"^SRS-[A-Z][A-Z0-9-]*-\d{3,}$")
PLACEHOLDER_RE = re.compile(r"(?:3\.2\.X|3\.3\.X|VX\.X|20\d{2}年XX月|\(CSCI能力\)|接口的唯一标识符)")
VAGUE_TERM_RE = re.compile(r"(?:友好|高效|尽量|适当|必要时|快速地|灵活方便)")

# Defined in the template style sheet but never used in the document body.
FORBIDDEN_CONTENT_STYLES = {"DB表头", "DB表正文"}

TABLE_HEADER_STYLE = "145表头"
TABLE_BODY_STYLE = "145表正文"
FIGURE_STYLE = "145图样式"

REQUIRED_H1 = [
    "概述",
    "使用与技术指标要求",
    "总体架构",
    "分系统设计",
    "系统集成设计",
    "关键技术分析",
    "效能分析",
    "工程组织管理",
    "初步工作计划",
]

GJB_DOCUMENT_TYPES = {
    "SDP": ("软件开发计划", "A"),
    "SIP": ("软件安装计划", "B"),
    "STrP": ("软件移交计划", "C"),
    "STP": ("软件测试计划", "D"),
    "OCD": ("运行方案说明", "E"),
    "SSS": ("系统/子系统规格说明", "F"),
    "IRS": ("接口需求规格说明", "G"),
    "SSDD": ("系统/子系统设计说明", "H"),
    "IDD": ("接口设计说明", "I"),
    "SRS": ("软件需求规格说明", "J"),
    "SDD": ("软件设计说明", "K"),
    "DBDD": ("数据库设计说明", "L"),
    "STD": ("软件测试说明", "M"),
    "STR": ("软件测试报告", "N"),
    "SPS": ("软件产品规格说明", "O"),
    "SVD": ("软件版本说明", "P"),
    "SUM": ("软件用户手册", "Q"),
    "CPM": ("计算机编程手册", "R"),
    "FSM": ("固件保障手册", "S"),
    "SDSR": ("软件研制总结报告", "T"),
}

STRICT_H2 = {
    "概述": {"任务依据", "编制目的", "指导原则", "建设目标", "主要工作", "与其他项目关系", "名词术语", "缩略语"},
    "使用与技术指标要求": {"使用要求", "技术指标要求"},
    "总体架构": {"体系架构设计", "逻辑架构设计", "系统架构设计", "功能架构设计", "数据架构设计", "部署架构设计", "数据交换设计", "关键性能指标设计", "组织运用模式", "技术体制"},
    "工程组织管理": {"组织机构", "考核验收", "安全保密"},
}

LEGACY_TERMS = ["ZBZQ", "XX数据中台", "联合XX数据资源体系", "XX数据融合治理"]
FORBIDDEN_LOCAL_PATH_TOKENS = [
    "/" + "Users" + "/",
    "/" + "home" + "/",
    "C:" + "\\" + "Users" + "\\",
]

TYPED_NUMBER_RE = re.compile(r"^\d+(?:\.\d+)*[\s　]")
CAPTION_NUMBER_RE = re.compile(r"^(图|表)\s*\d+-\d+")
HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5"}

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_NS = {"w": W_NS, "r": R_NS, "rel": REL_NS}
W = f"{{{W_NS}}}"
R = f"{{{R_NS}}}"


def strip_heading_number(text: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\s+", "", text).strip()


@dataclass
class AuditResult:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    metrics: dict[str, int | str | bool] = field(default_factory=dict)
    findings: list[dict] = field(default_factory=list)
    checks: dict[str, str] = field(default_factory=dict)
    input_error: bool = False

    @property
    def ok(self) -> bool:
        return not self.errors

    def add_issue(self, message: str, severity: str = "error", rule_id: str = "DOCX.CHECK",
                  source: str = "implementation", clause: str = "", object_id: str = "",
                  location: str = "", status: str = "failed") -> None:
        self.findings.append(dict(rule_id=rule_id, severity=severity, source=source,
                                  clause=clause, object_id=object_id, location=location,
                                  message=message, status=status))
        (self.errors if severity == "error" else self.warnings).append(message)

    def add_error(self, message: str, **details) -> None:
        self.add_issue(message, severity="error", **details)

    def add_warning(self, message: str, **details) -> None:
        self.add_issue(message, severity="warning", **details)

    @property
    def review_status(self) -> str:
        if not self.ok:
            return "failed"
        if all(self.checks.get(key) == "passed" for key in ("source_coverage", "semantic_review", "visual_review")):
            return "ready"
        return "not_established"


TOC_STYLES = {f"toc {level}" for level in range(1, 10)}


def paragraph_texts(document: Document) -> Iterable[str]:
    """Visible content text, excluding TOC field results.

    TOC paragraphs are stale field caches regenerated by Word on open
    (w:updateFields), so legacy-term scans must not flag them.
    """
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name in TOC_STYLES:
            continue
        text = paragraph.text.strip()
        if text:
            yield text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        yield text


def is_cover_table(table) -> bool:
    """The template cover table uses cover styles in its first row."""
    cover_styles = {"编号密级", "文头字", "文件标识号", "文件名称", "单位名称"}
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name in cover_styles:
                return True
    return False


def xml_attr(element: ET.Element | None, namespace: str, local_name: str) -> str | None:
    if element is None:
        return None
    return element.get(f"{{{namespace}}}{local_name}")


def read_package_xml(package: zipfile.ZipFile, name: str, result: AuditResult) -> ET.Element | None:
    try:
        return ET.fromstring(package.read(name))
    except KeyError:
        result.add_error(f"DOCX package is missing required XML part: {name}")
    except ET.ParseError as exc:
        result.add_error(f"DOCX XML part is not parseable: {name}: {exc}")
    return None


def package_xml_text(root: ET.Element, xpath: str) -> str:
    return " ".join((node.text or "") for node in root.findall(xpath, XML_NS))


def document_has_toc_field(path: Path, result: AuditResult) -> bool:
    with zipfile.ZipFile(path) as package:
        document_root = read_package_xml(package, "word/document.xml", result)
        if document_root is None:
            return False
        instr_text = package_xml_text(document_root, ".//w:instrText")
        fld_simple_instr = " ".join(
            xml_attr(node, W_NS, "instr") or ""
            for node in document_root.findall(".//w:fldSimple", XML_NS)
        )
        return "TOC" in (instr_text + " " + fld_simple_instr).upper()


def toc_instruction(path: Path, result: AuditResult) -> str:
    with zipfile.ZipFile(path) as package:
        document_root = read_package_xml(package, "word/document.xml", result)
        if document_root is None:
            return ""
        values = [(node.text or "") for node in document_root.findall(".//w:instrText", XML_NS)]
        values.extend(xml_attr(node, W_NS, "instr") or "" for node in document_root.findall(".//w:fldSimple", XML_NS))
        return " ".join(values)


def relationship_targets(package: zipfile.ZipFile, result: AuditResult) -> dict[str, str]:
    root = read_package_xml(package, "word/_rels/document.xml.rels", result)
    if root is None:
        return {}
    targets: dict[str, str] = {}
    for rel in root.findall("rel:Relationship", XML_NS):
        rel_id = rel.get("Id")
        target = rel.get("Target")
        if rel_id and target:
            targets[rel_id] = target
    return targets


def footer_part_name(target: str | None) -> str | None:
    if not target:
        return None
    return target if target.startswith("word/") else f"word/{target}"


def footer_has_page_field(package: zipfile.ZipFile, part_name: str | None, result: AuditResult) -> bool:
    if not part_name:
        return False
    root = read_package_xml(package, part_name, result)
    if root is None:
        return False
    instr_text = " ".join((node.text or "") for node in root.findall(".//w:instrText", XML_NS))
    instr_text += " " + " ".join(node.get(W + "instr", "") for node in root.findall(".//w:fldSimple", XML_NS))
    return bool(re.search(r"\bPAGE\b", instr_text, re.I))


def footer_has_centered_page_field(package: zipfile.ZipFile, part_name: str | None, result: AuditResult) -> bool:
    if not part_name:
        return False
    root = read_package_xml(package, part_name, result)
    if root is None:
        return False
    for paragraph in root.findall(".//w:p", XML_NS):
        instr_text = " ".join((node.text or "") for node in paragraph.findall(".//w:instrText", XML_NS))
        if "PAGE" not in instr_text.upper():
            continue
        jc = paragraph.find("./w:pPr/w:jc", XML_NS)
        if xml_attr(jc, W_NS, "val") == "center":
            return True
    return False


def footer_has_dash_page_dash(package: zipfile.ZipFile, part_name: str | None, result: AuditResult) -> bool:
    if not part_name:
        return False
    root = read_package_xml(package, part_name, result)
    if root is None:
        return False
    has_page = footer_has_page_field(package, part_name, result)
    text_runs = [node.text or "" for node in root.findall(".//w:t", XML_NS)]
    dash_count = sum(text.count("—") + text.count("―") for text in text_runs)
    return has_page and dash_count >= 2


def audit_page_numbering(path: Path, result: AuditResult, document_type: str | None = None) -> None:
    """Audit PAGE fields and sections for the selected template profile."""
    DOCUMENT_PROFILES = load_document_profiles()
    with zipfile.ZipFile(path) as package:
        document_root = read_package_xml(package, "word/document.xml", result)
        if document_root is None:
            return
        relationships = relationship_targets(package, result)
        sections = document_root.findall(".//w:sectPr", XML_NS)
        profile = "SRS" if document_type == "SRS" or (document_type is None and len(sections) == 3) else "overall"
        profile_key = "SRS-P09" if profile == "SRS" else "overall-technical-solution"
        expected_count = int(DOCUMENT_PROFILES[profile_key]["sections"])
        result.metrics["section_count"] = len(sections)
        if len(sections) != expected_count:
            result.add_error(f"{profile} profile expects {expected_count} sections, got {len(sections)}")
            return

        for index, section in enumerate(sections, start=1):
            page_size = section.find("w:pgSz", XML_NS)
            width = xml_attr(page_size, W_NS, "w")
            height = xml_attr(page_size, W_NS, "h")
            if width != "11906" or height != "16838":
                result.add_error(f"Section {index} is not A4 11906 x 16838 twips: {width} x {height}")
            if profile == "SRS":
                page_margin = section.find("w:pgMar", XML_NS)
                expected = DOCUMENT_PROFILES["SRS-P09"]["section_margins_twips"][index - 1]
                actual = [xml_attr(page_margin, W_NS, key) for key in ["top", "bottom", "left", "right"]]
                if actual != [str(value) for value in expected]:
                    result.add_error(f"SRS section {index} margins differ from P-09 profile: expected {expected}, got {actual}")
                for key in ["header", "footer"]:
                    expected_distance = str(DOCUMENT_PROFILES["SRS-P09"][f"{key}_twips"])
                    actual_distance = xml_attr(page_margin, W_NS, key)
                    if actual_distance != expected_distance:
                        result.add_error(f"SRS section {index} {key} distance must be {expected_distance} twips, got {actual_distance}")

        if profile == "SRS":
            cover_section, front_section, body_section = sections
            cover_number = cover_section.find("w:pgNumType", XML_NS)
            if xml_attr(cover_number, W_NS, "fmt") != "lowerRoman" or xml_attr(cover_number, W_NS, "start") != "1":
                result.add_error("SRS cover must begin lowerRoman numbering at i", rule_id="GJB.COVER_PAGE", source="standard", clause="4.3.7.1")
            cover_footers = {xml_attr(ref, W_NS, "type"): footer_part_name(relationships.get(xml_attr(ref, R_NS, "id"))) for ref in cover_section.findall("w:footerReference", XML_NS)}
            active_type = "first" if cover_section.find("w:titlePg", XML_NS) is not None else "default"
            if not footer_has_page_field(package, cover_footers.get(active_type), result):
                result.add_error("SRS cover must have a visible native PAGE field", rule_id="GJB.COVER_PAGE", source="standard", clause="4.3.7.1")
        else:
            front_section, body_section = sections
            if front_section.find("w:titlePg", XML_NS) is None:
                result.add_error("Front-matter section must preserve titlePg")

        front_page_numbering = front_section.find("w:pgNumType", XML_NS)
        expected_front = DOCUMENT_PROFILES[profile_key]["front_page_format"]
        expected_start = None if profile == "SRS" else "1"
        if xml_attr(front_page_numbering, W_NS, "fmt") != expected_front or xml_attr(front_page_numbering, W_NS, "start") != expected_start:
            result.add_error(f"Front-matter section must use {expected_front}; numbering must {'continue from the cover' if profile == 'SRS' else 'start at 1'}", source="template", rule_id="PROFILE.PAGE_FORMAT")
        if expected_front != "lowerRoman":
            result.add_warning("The selected legacy template uses upper Roman front matter, a deviation from GJB 438C 4.3.7.1", source="standard", clause="4.3.7.1", rule_id="GJB.FRONT_ROMAN")
        if profile == "SRS" and front_section.find("w:titlePg", XML_NS) is not None:
            result.add_error("SRS front matter must display the page number on its first page", source="template", rule_id="PROFILE.FRONT_FIRST_PAGE")

        front_footer_parts = [footer_part_name(relationships.get(xml_attr(ref, R_NS, "id"))) for ref in front_section.findall("w:footerReference", XML_NS)]
        if not front_footer_parts or not any(footer_has_page_field(package, part, result) for part in front_footer_parts):
            result.add_error("Front-matter footer must contain a native PAGE field")

        body_page_numbering = body_section.find("w:pgNumType", XML_NS)
        body_fmt = xml_attr(body_page_numbering, W_NS, "fmt")
        body_start = xml_attr(body_page_numbering, W_NS, "start")
        if profile == "SRS" and (body_fmt not in {None, "decimal"} or body_start != "1"):
            result.add_error("SRS body section must use Arabic page numbers starting at 1")
        if not body_section.findall("w:headerReference", XML_NS):
            result.add_error("Body section must preserve header references")
        body_footer_parts = [footer_part_name(relationships.get(xml_attr(ref, R_NS, "id"))) for ref in body_section.findall("w:footerReference", XML_NS)]
        body_page_footers = [part for part in body_footer_parts if footer_has_page_field(package, part, result)]
        if not body_page_footers:
            result.add_error("Body footer must contain a native PAGE field")
        elif not any(footer_has_dash_page_dash(package, part, result) for part in body_page_footers):
            result.add_error("Body footer must preserve the dash-PAGE-dash pattern")

        settings_root = read_package_xml(package, "word/settings.xml", result)
        update_fields = settings_root.find("w:updateFields", XML_NS) if settings_root is not None else None
        if update_fields is None or xml_attr(update_fields, W_NS, "val") not in {None, "true", "1", "on"}:
            result.add_warning("No automatic field-refresh request is saved; verify final field caches and rendered pages independently (editors may remove w:updateFields after refresh)")


def audit_gjb_composition(path: Path, document: Document, result: AuditResult, strict_gjb_composition: bool, document_type: str | None) -> None:
    def add_composition_issue(message: str) -> None:
        if strict_gjb_composition:
            result.add_error(message)
        else:
            result.add_warning(message)

    if document_type == "SRS":
        cover_text = "\n".join(p.text for p in document.paragraphs[:24])
        cover_compact = re.sub(r"\s+", "", cover_text)
        for label in ["密", "版", "阶段标注", "编", "审核", "标", "批准"]:
            if label not in cover_compact:
                result.add_error(f"SRS paragraph cover is missing required label containing {label!r}")
    elif not document.tables or not any(is_cover_table(table) for table in document.tables):
        result.add_error("Overall-solution profile must preserve its cover table")

    if not document_has_toc_field(path, result):
        result.add_error("Document must preserve a real Word TOC field instead of static TOC text")
    elif document_type == "SRS":
        instruction = toc_instruction(path, result)
        for token in ['1-3', r'\h', r'\z', r'\u']:
            if token.lower() not in instruction.lower():
                result.add_error(f"SRS TOC field is missing required switch/token {token}")

    full_visible_text = "\n".join(paragraph_texts(document))
    if not any(token in full_visible_text for token in ["修改页", "修改记录", "修改历史", "修订记录", "变更记录"]):
        add_composition_issue("GJB 438C composition expects a modification page/history; none was found")

    heading1_count = sum(
        1 for paragraph in document.paragraphs
        if paragraph.text.strip() and paragraph.style and paragraph.style.name == "Heading 1"
    )
    if heading1_count == 0:
        result.add_error("Document body must contain at least one Heading 1 body chapter/clause")


def table_rows(table) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def split_methods(value: str) -> set[str]:
    return {item.strip() for item in re.split(r"[、,，;/；]+", value) if item.strip()}


def table_has_repeat_header(table) -> bool:
    return bool(table.rows and table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))


def row_prevents_split(row) -> bool:
    return bool(row._tr.xpath("./w:trPr/w:cantSplit"))


def table_has_visible_borders(table) -> bool:
    def edges(element):
        names = {node.tag.rsplit("}", 1)[-1] for node in element if node.get(qn("w:val")) not in {None, "nil", "none"}}
        if "start" in names:
            names.add("left")
        if "end" in names:
            names.add("right")
        return names

    borders = table._tbl.xpath("./w:tblPr/w:tblBorders")
    required = {"top", "left", "bottom", "right", "insideH", "insideV"}
    if borders and required.issubset(edges(borders[0])):
        return True
    # Some editors materialize the same visible grid as per-cell borders.
    # Require all four edges of every cell, rather than assuming any one cell
    # or a named table style proves the entire table has a visible grid.
    cells = table._tbl.xpath("./w:tr/w:tc")
    for cell in cells:
        cell_borders = cell.find("./" + qn("w:tcPr") + "/" + qn("w:tcBorders"))
        if cell_borders is None or not {"top", "left", "bottom", "right"}.issubset(edges(cell_borders)):
            return False
    return bool(cells)


def table_has_fixed_layout(table) -> bool:
    nodes = table._tbl.xpath("./w:tblPr/w:tblLayout")
    return bool(nodes and nodes[0].get(qn("w:type")) == "fixed")


def table_geometry_issue(table) -> str | None:
    tbl_w_nodes = table._tbl.xpath("./w:tblPr/w:tblW")
    if not tbl_w_nodes or tbl_w_nodes[0].get(qn("w:type")) != "dxa":
        return "tblW is not an explicit dxa width"
    try:
        total = int(tbl_w_nodes[0].get(qn("w:w")) or "0")
        grid = [int(node.get(qn("w:w")) or "0") for node in table._tbl.xpath("./w:tblGrid/w:gridCol")]
    except ValueError:
        return "table width contains a non-numeric value"
    if total <= 0 or not grid or sum(grid) != total:
        return f"tblW/grid mismatch: tblW={total}, grid={grid}"
    for row_number, row in enumerate(table.rows, start=1):
        cell_widths = []
        for cell in row.cells:
            nodes = cell._tc.xpath("./w:tcPr/w:tcW")
            if not nodes or nodes[0].get(qn("w:type")) != "dxa":
                return f"row {row_number} lacks explicit dxa tcW"
            try:
                cell_widths.append(int(nodes[0].get(qn("w:w")) or "0"))
            except ValueError:
                return f"row {row_number} contains a non-numeric tcW"
        if cell_widths != grid:
            return f"row {row_number} tcW differs from tblGrid: tcW={cell_widths}, grid={grid}"
    return None


def audit_package_cleanliness(path: Path, result: AuditResult) -> None:
    import posixpath
    revision_tags = {"ins", "del", "moveFrom", "moveTo", "moveFromRangeStart", "moveToRangeStart",
                     "pPrChange", "rPrChange", "tblPrChange", "trPrChange", "tcPrChange", "sectPrChange", "numberingChange"}
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        bookmarks, references, instructions = set(), [], []
        for name in sorted(n for n in names if n.startswith("word/") and n.endswith(".xml")):
            root = read_package_xml(package, name, result)
            if root is None:
                continue
            if name.startswith("word/comments"):
                result.add_error("DOCX contains review comments", rule_id="DOCX.COMMENTS", source="review", location=name)
            if any(node.tag.startswith(W) and node.tag[len(W):] in revision_tags for node in root.iter()):
                result.add_error("DOCX contains tracked revisions", rule_id="DOCX.REVISIONS", source="review", location=name)
            text = package_xml_text(root, ".//w:t")
            if "计划阶段-06-软件需求规格说明" in text:
                result.add_error("Legacy P-09 header remains", rule_id="PROFILE.HEADER", source="template", location=name)
            for token in FORBIDDEN_LOCAL_PATH_TOKENS:
                if token in text:
                    result.add_error("Local workstation path remains in document part", rule_id="DOCX.LOCAL_PATH", source="review", location=name)
            starts = root.findall(".//w:bookmarkStart", XML_NS)
            ends = {xml_attr(n, W_NS, "id") for n in root.findall(".//w:bookmarkEnd", XML_NS)}
            for node in starts:
                bname = xml_attr(node, W_NS, "name") or ""
                if bname in bookmarks:
                    result.add_error(f"Duplicate bookmark name: {bname}", rule_id="DOCX.BOOKMARK_DUPLICATE", location=name)
                bookmarks.add(bname)
                if xml_attr(node, W_NS, "id") not in ends:
                    result.add_error(f"Bookmark has no matching end: {bname}", rule_id="DOCX.BOOKMARK_END", location=name)
            # Concatenate split instruction runs within each complex field.
            stack = []
            for node in root.iter():
                if node.tag == W + "fldSimple":
                    instructions.append((name, node.get(W + "instr", "")))
                elif node.tag == W + "fldChar":
                    kind = node.get(W + "fldCharType")
                    if kind == "begin":
                        stack.append([])
                    elif kind == "end" and stack:
                        instructions.append((name, "".join(stack.pop())))
                elif node.tag == W + "instrText":
                    if stack:
                        stack[-1].append(node.text or "")
                    else:
                        instructions.append((name, node.text or ""))
            for parts in stack:
                instructions.append((name, "".join(parts)))
            references.extend((name, n.get(W + "anchor")) for n in root.findall(".//w:hyperlink", XML_NS) if n.get(W + "anchor"))
        for name, instruction in instructions:
            for match in re.finditer(r'\b(?:PAGEREF|REF)\s+(?:"([^"]+)"|([^\s\\]+))', instruction, flags=re.I):
                references.append((name, match.group(1) or match.group(2)))
        for name, target in references:
            if target not in bookmarks:
                result.add_error(f"Cross-reference targets missing bookmark: {target}", rule_id="DOCX.DANGLING_REFERENCE", location=name)
        for name in sorted(n for n in names if n.endswith(".rels")):
            root = ET.fromstring(package.read(name))
            base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
            for rel in root:
                target = rel.get("Target", "")
                if rel.get("TargetMode") == "External":
                    if target.lower().startswith("file:") or any(t in target for t in FORBIDDEN_LOCAL_PATH_TOKENS):
                        result.add_error("Relationship links to a local workstation file", rule_id="DOCX.LOCAL_RELATIONSHIP", source="review", location=name)
                else:
                    from urllib.parse import unquote
                    resolved = posixpath.normpath(posixpath.join(base, unquote(target))).lstrip("/")
                    if resolved not in names:
                        result.add_error(f"Relationship target is missing: {target}", rule_id="DOCX.BROKEN_RELATIONSHIP", location=name)
        result.metrics["bookmark_count"] = len(bookmarks)
        result.metrics["seq_field_count"] = sum(bool(re.search(r"\bSEQ\s+", i, re.I)) for _, i in instructions)
        if not any(n.startswith("_Toc") for n in bookmarks):
            result.add_warning("No cached TOC bookmarks found; update fields and inspect the saved contents pages", rule_id="DOCX.TOC_REFRESH", source="review")


def audit_srs_semantics(path: Path, document: Document, result: AuditResult, strict: bool,
                        template_profile: str = "auto", content: dict | None = None) -> None:
    from srs_docx_audit import audit_srs
    audit_srs(path, document, result, strict, template_profile, content)


def audit_docx(path: Path, strict_secondary: bool = False, strict_tables: bool = False,
               strict_overall: bool = False, strict_gjb_composition: bool = False,
               document_type: str | None = None, strict_srs: bool = False,
               mode: str = "draft", template_profile: str = "auto", content: dict | None = None) -> AuditResult:
    result = AuditResult()
    try:
        if not path.is_file() or not zipfile.is_zipfile(path):
            raise ValueError(f"Not a readable DOCX ZIP package: {path}")
        document = Document(path)
        result.metrics["artifact_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        if strict_srs:
            document_type, mode = "SRS", "review"
        if document_type is None:
            titles = [strip_heading_number(p.text.strip()) for p in document.paragraphs
                      if p.style and p.style.name == "Heading 1"]
            if titles[:6] == SRS_REQUIRED_H1:
                document_type = "SRS"
        if document_type == "SRS":
            from srs_docx_audit import selected_profile
            profile = selected_profile(document, template_profile)
        else:
            profile = "overall-technical-solution"
        if profile != "standard":
            audit_page_numbering(path, result, document_type=document_type)
            styles = {style.name for style in document.styles}
            required = SRS_REQUIRED_STYLES if document_type == "SRS" else REQUIRED_STYLES
            missing = sorted(required - styles)
            if missing:
                result.add_error(f"Missing selected-template styles: {missing}", rule_id="PROFILE.STYLES", source="template")
            audit_gjb_composition(path, document, result, strict_gjb_composition or mode == "review", document_type)
        else:
            audit_standard_composition(document, result, mode == "review")
        audit_package_cleanliness(path, result)
        full_text = "\n".join(paragraph_texts(document))
        for token in FORBIDDEN_LOCAL_PATH_TOKENS:
            if token in full_text:
                result.add_error(f"Local environment path remains in document: {token}", rule_id="DOCX.LOCAL_PATH", source="review")
        for term in LEGACY_TERMS:
            if term in full_text:
                result.add_error(f"Legacy template term remains: {term}", rule_id="PROFILE.PLACEHOLDER", source="template")
        if document_type == "SRS":
            audit_srs_semantics(path, document, result, mode == "review", profile, content)
            return result

        # Preserve the existing overall-solution/other-document audit surface.
        h1_seen, current_primary = [], None
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text or not paragraph.style:
                continue
            style_name = paragraph.style.name
            if style_name in TOC_STYLES:
                continue
            if style_name in HEADING_STYLES:
                if TYPED_NUMBER_RE.match(text):
                    result.add_error(f"Heading has a typed number: {text[:60]}", rule_id="PROFILE.DOUBLE_NUMBER", source="template")
                normalized = strip_heading_number(text)
                if style_name == "Heading 1":
                    h1_seen.append(normalized)
                    current_primary = normalized
                elif style_name == "Heading 2" and strict_secondary and current_primary in STRICT_H2:
                    if normalized not in STRICT_H2[current_primary]:
                        result.add_error(f"Second-level heading outside selected template: {text}", source="template")
            elif style_name in FORBIDDEN_CONTENT_STYLES or style_name == "Normal":
                result.add_error(f"Body paragraph uses incompatible template style {style_name}: {text[:60]}", source="template")
        if (strict_overall or strict_secondary) and h1_seen != REQUIRED_H1:
            result.add_error(f"Overall template chapter order mismatch: {h1_seen}", source="template")
        for i, table in enumerate(document.tables, start=1):
            if is_cover_table(table):
                continue
            for j, row in enumerate(table.rows):
                expected = TABLE_HEADER_STYLE if j == 0 else TABLE_BODY_STYLE
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if paragraph.text.strip() and paragraph.style.name != expected:
                            result.add_issue(f"Table {i} row {j+1} style {paragraph.style.name}, expected {expected}",
                                severity="error" if strict_tables else "warning", source="template", rule_id="PROFILE.TABLE_STYLE")
        result.checks["machine"] = "passed" if result.ok else "failed"
    except Exception as exc:
        # A malformed package/profile is an input/runtime error, not a Python
        # traceback on a documented JSON interface.
        result.input_error = True
        result.add_error(f"Cannot audit input: {type(exc).__name__}: {exc}", rule_id="INPUT.DOCX", source="input", location=str(path))
    return result


def audit_standard_composition(document: Document, result: AuditResult, strict: bool) -> None:
    from srs_docx_audit import heading_level
    def issue(message, clause):
        result.add_issue(message, severity="error" if strict else "warning",
                         rule_id="GJB.COMPOSITION", source="standard", clause=clause)
    text = "\n".join(paragraph_texts(document))
    if not any(word in text for word in ("修改历史", "修改记录", "修改页", "修订记录")):
        issue("No modification-history information identified", "4.3.3")
    # Word fields are the bundled template mechanism, not the only standard-
    # conforming way to represent an approved, paginated table of contents.
    if not re.search(r"目\s*录", text):
        result.add_warning("Verify that the document has a complete contents list", rule_id="GJB.CONTENTS_REVIEW", source="semantic", clause="4.3.4", status="needs_review")
    section_index, body_section_index = 0, None
    for paragraph in document.paragraphs:
        if heading_level(paragraph) == 1 and body_section_index is None:
            body_section_index = section_index
        if paragraph._p.xpath("./w:pPr/w:sectPr"):
            section_index += 1
    for i, section in enumerate(document.sections):
        number = section._sectPr.find(qn("w:pgNumType"))
        fmt = number.get(qn("w:fmt"), "decimal") if number is not None else "decimal"
        if body_section_index is not None and i < body_section_index:
            # An unnumbered cover has no visible PAGE field and may use titlePg.
            footers = [section.footer, section.first_page_footer]
            numbered = any("PAGE" in f._element.xml for f in footers)
            if numbered and fmt != "lowerRoman":
                issue(f"Front-matter section {i+1} uses {fmt}, expected lowerRoman", "4.3.7.1")
        elif fmt not in {"decimal", "decimalZero"}:
            issue(f"Body/appendix section {i+1} uses {fmt}, expected Arabic numbering", "4.3.7.1")


def parse_frontmatter(skill_md: Path) -> dict[str, str]:
    text = skill_md.read_text(encoding="utf-8")
    if not text.startswith("---\n"):
        raise ValueError("SKILL.md must start with YAML frontmatter")
    end = text.find("\n---", 4)
    if end < 0:
        raise ValueError("SKILL.md frontmatter is not closed")
    metadata: dict[str, str] = {}
    for line in text[4:end].splitlines():
        if not line.strip() or line.lstrip().startswith("#") or ":" not in line:
            continue
        key, value = line.split(":", 1)
        metadata[key.strip()] = value.strip().strip('"\'')
    return metadata


def iter_text_files(root: Path) -> Iterable[Path]:
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in {".md", ".py", ".yml", ".yaml", ".toml", ".txt"}:
            yield path


def validate_skill(skill_dir: Path) -> AuditResult:
    skill_dir = skill_dir.resolve()
    result = AuditResult()
    required_paths = [
        skill_dir / "SKILL.md",
        skill_dir / "agents" / "openai.yaml",
        skill_dir / "assets" / "1-2总体技术方案-模板参考.docx",
        skill_dir / "assets" / "P-09-软件需求规格说明-438C模板.docx",
        skill_dir / "references" / "template-structure.md",
        skill_dir / "references" / "gjb438c-2021-requirements.md",
        skill_dir / "references" / "format-lock.md",
        skill_dir / "references" / "writing-prompts.md",
        skill_dir / "references" / "srs-appendix-j.md",
        skill_dir / "references" / "srs-p09-format-lock.md",
        skill_dir / "references" / "srs-requirement-model.md",
        skill_dir / "references" / "srs-review-gate.md",
        skill_dir / "references" / "document-profiles.json",
        skill_dir / "references" / "srs-content-schema.json",
        skill_dir / "scripts" / "apply_gjb438c_template.py",
        skill_dir / "scripts" / "audit_gjb438c_docx.py",
        skill_dir / "scripts" / "srs_model.py",
        skill_dir / "scripts" / "srs_docx_audit.py",
    ]
    for path in required_paths:
        if not path.exists():
            result.add_error(f"Missing required skill file: {path}")

    if (skill_dir / "README.md").exists():
        result.add_error("Skill body must not contain README.md; keep repository docs at repo root.")

    skill_md = skill_dir / "SKILL.md"
    if skill_md.exists():
        try:
            metadata = parse_frontmatter(skill_md)
            if metadata.get("name") != skill_dir.name:
                result.add_error(f"Skill name {metadata.get('name')!r} does not match folder name {skill_dir.name!r}")
            if not metadata.get("description"):
                result.add_error("SKILL.md description is required")
        except Exception as exc:
            result.add_error(f"Invalid SKILL.md frontmatter: {exc}")

    template = skill_dir / "assets" / "1-2总体技术方案-模板参考.docx"
    if template.exists() and not zipfile.is_zipfile(template):
        result.add_error(f"Template asset is not a valid DOCX package: {template}")
    srs_template = skill_dir / "assets" / "P-09-软件需求规格说明-438C模板.docx"
    if srs_template.exists() and not zipfile.is_zipfile(srs_template):
        result.add_error(f"SRS template asset is not a valid DOCX package: {srs_template}")

    for json_name in ["document-profiles.json", "srs-content-schema.json"]:
        json_path = skill_dir / "references" / json_name
        if json_path.exists():
            try:
                value = json.loads(json_path.read_text(encoding="utf-8"))
                if json_name == "srs-content-schema.json":
                    from jsonschema import Draft202012Validator
                    Draft202012Validator.check_schema(value)
                elif not isinstance(value, dict) or not all(key in value for key in ("SRS-P09", "overall-technical-solution")):
                    raise ValueError("Missing document profiles")
                elif value["SRS-P09"].get("front_page_format") != "lowerRoman":
                    raise ValueError("SRS profile must implement GJB 4.3.7.1 lowerRoman front matter")
            except Exception as exc:
                result.add_error(f"Invalid JSON reference {json_path}: {exc}")

    for script_name in ["apply_gjb438c_template.py", "audit_gjb438c_docx.py", "srs_model.py", "srs_docx_audit.py"]:
        script_path = skill_dir / "scripts" / script_name
        if script_path.exists():
            try:
                compile(script_path.read_text(encoding="utf-8"), str(script_path), "exec")
            except SyntaxError as exc:
                result.add_error(f"Python script does not compile: {script_path}: {exc}")

    for path in iter_text_files(skill_dir):
        text = path.read_text(encoding="utf-8", errors="ignore")
        for token in FORBIDDEN_LOCAL_PATH_TOKENS:
            if token in text:
                result.add_error(f"Local environment path appears in {path}: {token}")
        if path.suffix == ".md":
            for target in re.findall(r"\]\(([^)]+)\)", text):
                if re.match(r"[a-z]+://|#", target) or " " in target:
                    continue
                local = target.split("#", 1)[0]
                if local and not (path.parent / local).exists():
                    result.add_error(f"Broken reference in {path.name}: {target}", rule_id="SKILL.BROKEN_LINK", location=str(path))

    return result


def print_result(result: AuditResult, as_json: bool) -> None:
    if as_json:
        print(json.dumps({"ok": result.ok, "errors": result.errors, "warnings": result.warnings,
                          "metrics": result.metrics, "findings": result.findings,
                          "checks": result.checks, "review_status": result.review_status}, ensure_ascii=False, indent=2))
        return
    if result.errors:
        print("ERRORS:")
        for item in result.errors:
            print(f"- {item}")
    if result.warnings:
        print("WARNINGS:")
        for item in result.warnings:
            print(f"- {item}")
    if result.ok:
        print("OK")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", nargs="?", type=Path, help="DOCX file to audit.")
    parser.add_argument("--skill-dir", type=Path, help="Skill directory to validate.")
    parser.add_argument("--validate-skill", action="store_true", help="Validate the skill bundle structure.")
    parser.add_argument("--strict-secondary", action="store_true", help="Require fixed second-level headings where the template defines them.")
    parser.add_argument("--strict-overall", action="store_true", help="Require the overall technical solution's fixed nine first-level chapters.")
    parser.add_argument("--strict-gjb-composition", action="store_true", help="Treat missing GJB composition elements such as modification history as errors.")
    parser.add_argument("--strict-tables", action="store_true", help="Treat table style mismatches as errors.")
    parser.add_argument("--document-type", choices=["SRS", "overall-technical-solution"], help="Select a document-specific template/audit profile.")
    parser.add_argument("--strict-srs", action="store_true", help="Alias for SRS review-mode machine checks; does not certify semantic/visual readiness.")
    parser.add_argument("--mode", choices=["draft", "review"], default="draft", help="SRS missing-information policy (default: draft).")
    parser.add_argument("--template-profile", choices=["auto", "standard", "SRS-P09"], default="auto", help="Keep Appendix J content checks separate from selected-template checks.")
    parser.add_argument("--content-json", type=Path, help="Independent source/requirement model to validate and compare to the saved SRS.")
    parser.add_argument("--json", action="store_true", help="Print machine-readable audit output.")
    return parser.parse_args()


def merge_results(results: list[AuditResult]) -> AuditResult:
    merged = AuditResult()
    for result in results:
        merged.errors.extend(result.errors)
        merged.warnings.extend(result.warnings)
        merged.metrics.update(result.metrics)
        merged.findings.extend(result.findings)
        merged.checks.update(result.checks)
        merged.input_error = merged.input_error or result.input_error
    return merged


def main() -> int:
    # Keep Chinese findings and paths intact under redirected Windows output.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    args = parse_args()
    results: list[AuditResult] = []
    try:
        if args.validate_skill:
            if not args.skill_dir:
                raise ValueError("--validate-skill requires --skill-dir")
            results.append(validate_skill(args.skill_dir))
        content = json.loads(args.content_json.read_text(encoding="utf-8")) if args.content_json else None
        if args.docx:
            results.append(audit_docx(
                args.docx,
                args.strict_secondary,
                args.strict_tables,
                strict_overall=args.strict_overall,
                strict_gjb_composition=args.strict_gjb_composition,
                document_type=args.document_type,
                strict_srs=args.strict_srs,
                mode=args.mode, template_profile=args.template_profile, content=content,
            ))
        if not results:
            raise ValueError("Provide a DOCX path and/or --validate-skill --skill-dir")
    except Exception as exc:
        failure = AuditResult(input_error=True)
        failure.add_error(f"{type(exc).__name__}: {exc}", rule_id="INPUT.INVALID", source="input")
        results.append(failure)
    result = merge_results(results)
    print_result(result, args.json)
    return 2 if result.input_error else (0 if result.ok else 1)


if __name__ == "__main__":
    raise SystemExit(main())
