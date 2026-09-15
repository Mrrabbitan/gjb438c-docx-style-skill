"""End-to-end invariants between a source model and the final saved DOCX."""
from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "gjb438c-docx-style/scripts"))
sys.path.insert(0, str(ROOT / "tools"))
from apply_gjb438c_template import build_document
from audit_gjb438c_docx import audit_docx
from build_srs_fixture import fixture
from srs_model import SRS_CLAUSES

TEMPLATE = ROOT / "gjb438c-docx-style/assets/P-09-软件需求规格说明-438C模板.docx"


class SrsPipelineTests(unittest.TestCase):
    def run_model(self, model, mutate=None):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.docx"
            build_document(TEMPLATE, model, path, mode="review")
            if mutate:
                doc = Document(path); mutate(doc); doc.save(path)
            return audit_docx(path, document_type="SRS", strict_srs=True, content=model)

    def test_complete_model_passes_machine_checks_but_not_human_review(self):
        result = self.run_model(fixture())
        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.checks["source_coverage"], "passed")
        self.assertEqual(result.checks["visual_review"], "not_run")
        self.assertEqual(result.review_status, "not_established")

    def test_saved_requirement_attribute_change_is_detected(self):
        def mutate(doc):
            for table in doc.tables:
                for row in table.rows:
                    if row.cells[0].text == "适用条件":
                        row.cells[1].text = "未经来源允许的新条件"
                        return
        result = self.run_model(fixture(), mutate)
        self.assertIn("SRS.MODEL_ATTRIBUTE", {f["rule_id"] for f in result.findings})

    def test_toc_title_style_survives_editor_rebuild(self):
        def inspect(doc):
            style = doc.styles["TOC Heading"]
            self.assertEqual(style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(style.font.size.pt, 16)
            self.assertEqual(style.element.rPr.find(qn("w:color")).get(qn("w:val")), "000000")
            for name in ("TOC Heading", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "SRS表头", "SRS表正文", "SRS题注"):
                self.assertEqual(doc.styles[name].element.pPr.find(qn("w:ind")).get(qn("w:firstLineChars")), "0", name)
        self.assertTrue(self.run_model(fixture(), inspect).ok)

    def test_editor_cell_borders_and_removed_refresh_request_are_accepted(self):
        def editor_roundtrip(doc):
            flag = doc.settings.element.find(qn("w:updateFields"))
            doc.settings.element.remove(flag)
            for table in doc.tables:
                table_borders = table._tbl.tblPr.find(qn("w:tblBorders"))
                table._tbl.tblPr.remove(table_borders)
                for row in table.rows:
                    for cell in row.cells:
                        borders = OxmlElement("w:tcBorders")
                        for name in ("top", "start", "bottom", "end"):
                            edge = OxmlElement("w:" + name)
                            edge.set(qn("w:val"), "single")
                            borders.append(edge)
                        cell._tc.get_or_add_tcPr().append(borders)
        result = self.run_model(fixture(), editor_roundtrip)
        self.assertTrue(result.ok, result.errors)
        self.assertTrue(any("field-refresh" in warning for warning in result.warnings))

    def test_saved_forward_mapping_cannot_be_invented(self):
        def mutate(doc):
            for table in doc.tables:
                if "SRS需求编号" in [c.text for c in table.rows[0].cells] and "来源编号" in [c.text for c in table.rows[0].cells]:
                    row = table.add_row()
                    for cell, text in zip(row.cells, ["UNKNOWN-SOURCE", "软件", "UNKNOWN-REQ", "3.2", "落实"]):
                        cell.text = text
                    return
        result = self.run_model(fixture(), mutate)
        self.assertIn("SRS.MODEL_FORWARD", {f["rule_id"] for f in result.findings})

    def test_saved_qualification_criterion_change_is_detected(self):
        def mutate(doc):
            for table in doc.tables:
                headers = [c.text for c in table.rows[0].cells]
                if "通过准则" in headers:
                    table.rows[1].cells[headers.index("通过准则")].text = "随意通过"
                    return
        result = self.run_model(fixture(), mutate)
        self.assertIn("SRS.MODEL_QUALIFICATION", {f["rule_id"] for f in result.findings})

    def test_derived_requirements_accept_explicit_empty_allocated_baseline(self):
        model = fixture()
        model["source_requirements"], model["forwardTrace"], model["reverseTrace"] = [], [], []
        for req in model["requirements"]:
            req.pop("source_ids")
            req["derivation"] = {"basis": "DEMO-SSS 设计决策D1", "rationale": "该CSCI拆分边界需要此项行为约束。"}
            model["reverseTrace"].append(dict(requirement_id=req["id"], derivation_basis=req["derivation"]["basis"], methods=req["qualification_methods"]))
        result = self.run_model(model)
        self.assertTrue(result.ok, result.errors)

    def test_external_irs_targets_do_not_require_fictitious_local_requirements(self):
        model = fixture()
        model.update(requirements=[], states=[], capabilities=[], data=[])
        target = {"reference_id": "DEMO-IRS", "requirement_id": "EXTERNAL-7", "location": "3.2.1"}
        model["interfaces"] = [{"id": "IF-IRS", "name": "受控接口", "definition": "irs_reference", "irs_reference": {"reference_id": "DEMO-IRS", "location": "第3章"}}]
        model["source_requirements"] = [model["source_requirements"][2]]
        model["forwardTrace"] = [{"source": "UP-03", "external_target": copy.deepcopy(target), "disposition": "落实"}]
        model["reverseTrace"] = [{"external_target": copy.deepcopy(target), "source_ids": ["UP-03"], "methods": ["测试"]}]
        model["qualification"] = [{"external_target": target, "methods": ["测试"], "condition": "输入空编号文件", "evidence": "EV-IRS-7", "pass_criterion": "返回EMPTY_ID"}]
        model["tailoring"] = [{"clause": clause, "status": "not_applicable", "reason": "受控来源仅分配外部接口需求。", "basis": "DEMO-SSS第1章"} for clause in SRS_CLAUSES if clause not in {"3.3", "3.18"} and not clause.startswith("3.11.")]
        result = self.run_model(model)
        self.assertTrue(result.ok, result.errors)


if __name__ == "__main__":
    unittest.main()
