"""Independent saved-document mutations must not evade model comparison."""
from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "gjb438c-docx-style/scripts"))
sys.path.insert(0, str(ROOT / "tools"))
from build_srs_fixture import fixture
import apply_gjb438c_template as generator
from audit_gjb438c_docx import audit_docx
from docx import Document
from docx.oxml.ns import qn


class SavedModelConsistencyTests(unittest.TestCase):
    def run_model(self, model, mutate=None):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.docx"
            generator.build_document(generator.SRS_TEMPLATE, model, path, mode="review")
            if mutate:
                document = Document(path)
                mutate(document)
                document.save(path)
            return audit_docx(path, document_type="SRS", strict_srs=True, content=model)

    def replace_record_field(self, document, label, replacement):
        for table in document.tables:
            if [cell.text for cell in table.rows[0].cells] != ["字段", "内容"]:
                continue
            for row in table.rows[1:]:
                if row.cells[0].text == label:
                    row.cells[1].text = replacement
                    return
        self.fail(f"Generated fixture has no record field {label}")

    def closed_tbd_model(self):
        model = fixture()
        model["tbd"] = [{"id": "TBD-1", "issue": "记录范围已经确认", "affected_ids": [model["requirements"][0]["id"]], "owner": "需求负责人", "closure_condition": "批准范围基线", "status": "closed", "closure_evidence": "DEC-1"}]
        return model

    def test_closed_tbd_affected_objects_satisfy_impact_without_duplicate_prose(self):
        result = self.run_model(self.closed_tbd_model())
        self.assertTrue(result.ok, result.errors)

    def test_changed_tbd_owner_is_detected_against_model(self):
        model = self.closed_tbd_model()
        model["tbd"][0]["impact"] = "记录查询范围"
        result = self.run_model(model, lambda doc: self.replace_record_field(doc, "责任方", "错误责任方"))
        self.assertFalse(result.ok, "Changed TBD owner passed saved-document/model comparison")

    def test_changed_composite_data_definition_is_detected(self):
        result = self.run_model(fixture(), lambda doc: self.replace_record_field(doc, "数据元素", "损坏的字段定义"))
        self.assertFalse(result.ok, "Changed composite data definition passed model comparison")

    def test_composite_json_is_compared_structurally_and_keeps_nested_values(self):
        model = fixture()
        value = {"fields": [{"name": "record_id", "type": "string", "length": 32}], "version": 1}
        model["data"][0]["fields"] = value
        result = self.run_model(model, lambda doc: self.replace_record_field(doc, "数据元素", json.dumps(value, sort_keys=True)))
        self.assertTrue(result.ok, result.errors)
        changed = deepcopy(value)
        changed["fields"][0]["length"] = 64
        result = self.run_model(model, lambda doc: self.replace_record_field(doc, "数据元素", json.dumps(changed)))
        self.assertFalse(result.ok, "Changed nested data length passed structural model comparison")

    def test_changed_derivation_rationale_is_detected(self):
        model = fixture()
        requirement = model["requirements"][0]
        rid = requirement["id"]
        requirement.pop("source_ids")
        basis = "DEMO-SSS 第1章；DEC-1"
        requirement["derivation"] = {"basis": basis, "rationale": "CSCI分解需要该状态约束。"}
        model["forwardTrace"] = [row for row in model["forwardTrace"] if row["requirement_id"] != rid]
        for row in model["reverseTrace"]:
            if row["requirement_id"] == rid:
                row.pop("source_ids")
                row["derivation_basis"] = basis
        baseline = self.run_model(model)
        self.assertTrue(baseline.ok, baseline.errors)
        result = self.run_model(model, lambda doc: self.replace_record_field(doc, "派生依据", json.dumps({"basis": basis, "rationale": "未经来源支持的新派生理由"}, ensure_ascii=False)))
        self.assertFalse(result.ok, "Changed derivation rationale passed model comparison")

    def test_changed_controlled_reference_revision_is_detected(self):
        result = self.run_model(fixture(), lambda doc: self.replace_record_field(doc, "修订版", "错误版次"))
        self.assertFalse(result.ok, "Changed controlled reference revision passed model comparison")

    def test_qualification_conditions_and_criteria_must_remain_paired(self):
        model = fixture()
        first = model["qualification"][0]
        model["qualification"].append(dict(first, condition="条件 B", pass_criterion="判据 B", evidence="EV-B"))
        rid = first["requirement_id"]
        baseline = self.run_model(model)
        self.assertTrue(baseline.ok, baseline.errors)

        def swap_criteria(document):
            for table in document.tables:
                headers = [cell.text for cell in table.rows[0].cells]
                if "通过准则" not in headers:
                    continue
                rows = [row for row in table.rows[1:] if row.cells[0].text == rid]
                self.assertEqual(len(rows), 2)
                first_cell, second_cell = [row.cells[headers.index("通过准则")] for row in rows]
                first_cell.text, second_cell.text = second_cell.text, first_cell.text
                return
            self.fail("Generated fixture has no qualification matrix")

        result = self.run_model(model, swap_criteria)
        self.assertFalse(result.ok, "Cross-paired conditions and pass criteria passed model comparison")


class LayoutRegressionTests(unittest.TestCase):
    def test_modification_history_drops_only_empty_stationery_rows(self):
        document = Document(generator.SRS_TEMPLATE)
        generator.normalize_srs_styles(document)
        generator.clear_body_content(document)
        history = next(table for table in document.tables if [cell.text.strip() for cell in table.rows[0].cells] == ["版本", "修改原因", "修改内容", "修改人", "修改日期"])
        history.add_row()
        historical = history.add_row()
        for cell, value in zip(historical.cells, ["0.9", "历史变更", "保留已有变更事实", "记录人", "2026-09-01"]):
            cell.text = value
        generator.normalize_srs_headers_and_tables(document, {"version": "1.0", "change_reason": "本次变更", "change_content": "本次修订内容", "author": "编制人", "date": "2026-09-15"})
        records = [[cell.text.strip() for cell in row.cells] for row in history.rows]
        self.assertEqual(len(records), 3)
        self.assertEqual(records[1][2], "本次修订内容")
        self.assertEqual(records[2], ["0.9", "历史变更", "保留已有变更事实", "记录人", "2026-09-01"])
        self.assertFalse(history._tbl.xpath(".//w:trHeight"))

    def test_caption_fields_use_heading_number_and_independent_chapter_sequences(self):
        document = Document(generator.SRS_TEMPLATE)
        generator.normalize_srs_styles(document)
        generator.clear_body_content(document)
        generator.add_styled_paragraph(document, "第一章", "Heading 1")
        generator.add_caption_paragraph(document, "表 第一张", "SRS题注")
        generator.add_caption_paragraph(document, "表 第二张", "SRS题注")
        generator.add_caption_paragraph(document, "图 第一幅", "SRS题注")
        generator.add_styled_paragraph(document, "第二章", "Heading 1")
        generator.add_caption_paragraph(document, "表 另一章第一张", "SRS题注")
        captions = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "SRS题注"]
        self.assertEqual([paragraph.text for paragraph in captions], ["表 1-1 第一张", "表 1-2 第二张", "图 1-1 第一幅", "表 2-1 另一章第一张"])
        instructions = [[(node.text or "").strip() for node in paragraph._p.iter(qn("w:instrText"))] for paragraph in captions]
        self.assertTrue(all(row[0] == 'STYLEREF "Heading 1" \\n' for row in instructions))
        self.assertEqual([row[1] for row in instructions], ["SEQ Table_1 \\* ARABIC", "SEQ Table_1 \\* ARABIC", "SEQ Figure_1 \\* ARABIC", "SEQ Table_2 \\* ARABIC"])


if __name__ == "__main__":
    unittest.main()
