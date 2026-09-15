from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "gjb438c-docx-style" / "scripts"))
from audit_gjb438c_docx import AuditResult, audit_docx, audit_package_cleanliness
from srs_docx_audit import CHAPTERS, CLAUSES, RESOURCES, audit_srs, effective_numbering


def table(doc, fields, rows):
    obj = doc.add_table(rows=1, cols=len(fields))
    for cell, value in zip(obj.rows[0].cells, fields):
        cell.text = value
    for row in rows:
        for cell, value in zip(obj.add_row().cells, row):
            cell.text = value
    return obj


def simple_document(statement="软件应在收到停止指令后终止运行。"):
    doc = Document()
    elements = {}
    for i, title in enumerate(CHAPTERS, 1):
        elements[str(i)] = doc.add_heading(title, level=1)
        if i == 1:
            for text in ("标识", "系统概述", "文档概述"):
                doc.add_heading(text, level=2)
                doc.add_paragraph("公开虚构测试软件，版本1.0。")
        if i == 2:
            table(doc, ["文档编号", "标题", "编写单位", "修订版", "日期"], [["SRC", "虚构基线", "示例", "1.0", "2026-09-15"]])
        if i == 3:
            for j, clause_title in enumerate(CLAUSES, 1):
                elements[f"3.{j}"] = doc.add_heading(clause_title, level=2)
                if j == 1:
                    elements["requirement"] = table(doc, ["需求编号", "需求陈述", "合格性方法", "来源编号"], [["REQ-现有-01", statement, "测试", "SRC.1"]])
                if j == 3:
                    doc.add_heading("接口标识和接口图", level=3)
                if j == 11:
                    for k, resource in enumerate(RESOURCES, 1):
                        elements[f"3.11.{k}"] = doc.add_heading(resource, level=3)
        if i == 4:
            elements["qualification"] = table(doc, ["需求编号", "合格性方法", "验证条件", "计划证据", "通过准则"], [["REQ-现有-01", "测试", "正常运行", "EV-1", "停止后无新增输出"]])
        if i == 5:
            table(doc, ["来源编号", "SRS需求编号", "处置"], [["SRC.1", "REQ-现有-01", "落实"]])
    return doc, elements


def findings(doc):
    result = AuditResult()
    audit_srs(Path("memory.docx"), doc, result, strict=True, template_profile="standard")
    return result


class StructureTests(unittest.TestCase):
    def test_duplicate_document_requirement_is_error_even_in_draft(self):
        doc, _ = simple_document()
        table(doc, ["需求编号", "需求陈述"], [["REQ-现有-01", "软件应停止。"]])
        result = AuditResult()
        audit_srs(Path("memory.docx"), doc, result, strict=False, template_profile="standard")
        duplicates = [f for f in result.findings if f["rule_id"] == "SRS.REQUIREMENT_ID"]
        self.assertTrue(duplicates)
        self.assertEqual(duplicates[0]["severity"], "error")

    def test_valid_hierarchy_and_existing_id_are_not_rejected(self):
        doc, _ = simple_document()
        result = findings(doc)
        codes = {f["rule_id"] for f in result.findings}
        self.assertFalse(codes & {"SRS.CHAPTER_ORDER", "SRS.CLAUSE_PARENT", "SRS.MISSING_CLAUSE", "SRS.REQUIREMENT_ID"}, result.errors)
        self.assertEqual(result.checks["source_coverage"], "not_run")
        self.assertNotEqual(result.review_status, "ready")

    def test_resource_titles_in_wrong_parent_fail(self):
        doc, elements = simple_document()
        for n in range(1, 5):
            doc.element.body.insert(-1, elements[f"3.11.{n}"]._p)
        self.assertIn("SRS.CLAUSE_PARENT", {f["rule_id"] for f in findings(doc).findings})

    def test_matrix_moved_to_notes_fails(self):
        doc, elements = simple_document()
        doc.element.body.insert(-1, elements["qualification"]._tbl)
        self.assertIn("SRS.QUALIFICATION_LOCATION", {f["rule_id"] for f in findings(doc).findings})

    def test_parent_tailoring_covers_all_resource_children(self):
        doc, elements = simple_document()
        for n in range(1, 5):
            element = elements[f"3.11.{n}"]._p
            element.getparent().remove(element)
        paragraph = OxmlElement("w:p")
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "本条无内容。原因：基线第1章明确资源选择留待设计。"
        run.append(text); paragraph.append(run)
        elements["3.11"]._p.addnext(paragraph)
        result = findings(doc)
        self.assertFalse(any(f["rule_id"] == "SRS.MISSING_CLAUSE" and f["clause"].startswith("J.3.11.") for f in result.findings), result.errors)

    def test_empty_tbd_header_does_not_mask_unknown_fact(self):
        doc, _ = simple_document()
        doc.add_paragraph("待确认：软件恢复时间。")
        table(doc, ["TBD编号", "未决事项", "影响范围", "责任方", "关闭条件", "状态"], [])
        self.assertIn("SRS.UNRESOLVED_FACT", {f["rule_id"] for f in findings(doc).findings})

    def test_semantic_keyword_is_review_hint_not_blocker(self):
        for statement in ["响应正常。", "系统应在必要时（负载超过80%时）关闭非关键请求。"]:
            doc, _ = simple_document(statement)
            result = findings(doc)
            hints = [f for f in result.findings if f["rule_id"] == "SRS.SEMANTIC_CANDIDATE"]
            self.assertTrue(hints)
            self.assertTrue(all(f["severity"] == "warning" for f in hints))

    def test_interface_horizontal_and_vertical_records(self):
        for vertical in (False, True):
            doc, elements = simple_document()
            fields = ["接口编号", "CSCI端", "外部端", "方向"]
            values = ["IF-A", "CSCI", "FILE", "输入"]
            obj = table(doc, ["字段", "内容"], list(zip(fields, values))) if vertical else table(doc, fields, [values])
            elements["3.4"]._p.addnext(obj._tbl)
            self.assertEqual(findings(doc).metrics["interface_count"], 1)


class PackageTests(unittest.TestCase):
    def test_json_report_preserves_chinese_under_legacy_output_encoding(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "软件不存在.docx"
            process = subprocess.run([sys.executable, str(ROOT / "gjb438c-docx-style/scripts/audit_gjb438c_docx.py"), str(path), "--json"], capture_output=True, text=True, encoding="utf-8", env=dict(os.environ, PYTHONIOENCODING="cp1252"))
            self.assertEqual(process.returncode, 2, process.stderr)
            report = json.loads(process.stdout)
            self.assertFalse(report["ok"])
            self.assertIn("软件不存在", process.stdout)
            self.assertNotIn("Traceback", process.stderr)

    def test_invalid_docx_json_has_defined_exit(self):
        import zipfile
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.docx"
            with zipfile.ZipFile(path, "w") as z:
                z.writestr("word/document.xml", "<broken")
            process = subprocess.run([sys.executable, str(ROOT / "gjb438c-docx-style/scripts/audit_gjb438c_docx.py"), str(path), "--json"], text=True, capture_output=True, encoding="utf-8")
            self.assertEqual(process.returncode, 2)
            result = json.loads(process.stdout)
            self.assertFalse(result["ok"])
            self.assertNotIn("Traceback", process.stderr)

    def test_header_revisions_and_simple_reference_are_scanned(self):
        doc, _ = simple_document()
        paragraph = doc.sections[0].header.add_paragraph()
        insertion = OxmlElement("w:ins"); insertion.set(qn("w:id"), "1")
        paragraph._p.append(insertion)
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "REF MissingBookmark")
        paragraph._p.append(field)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"; doc.save(path)
            result = AuditResult(); audit_package_cleanliness(path, result)
        codes = {f["rule_id"] for f in result.findings}
        self.assertTrue({"DOCX.REVISIONS", "DOCX.DANGLING_REFERENCE"}.issubset(codes))

    def test_inherited_numbering_and_missing_definition(self):
        doc = Document()
        paragraph = doc.add_paragraph("条款", style="Heading 1")
        style = doc.styles["Heading 1"]
        ppr = style.element.get_or_add_pPr()
        numpr = OxmlElement("w:numPr"); numid = OxmlElement("w:numId"); level = OxmlElement("w:ilvl")
        numid.set(qn("w:val"), "1"); level.set(qn("w:val"), "0")
        numpr.extend([numid, level]); ppr.append(numpr)
        value, error = effective_numbering(paragraph, doc)
        self.assertIsNone(error)
        self.assertEqual(value[0], 1)
        numid.set(qn("w:val"), "999999")
        value, error = effective_numbering(paragraph, doc)
        self.assertIsNone(value)
        self.assertIn("no definition", error)


if __name__ == "__main__":
    unittest.main()
