"""Generated content must preserve facts and protect existing deliverables."""
from copy import deepcopy
from pathlib import Path
import json
import os
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_srs_model import valid_model, external_model, SCRIPTS
sys.path.insert(0, str(SCRIPTS))
import apply_gjb438c_template as generator
from docx import Document
from docx.oxml.ns import qn


class SrsGenerationTests(unittest.TestCase):
    def test_supported_capability_and_composite_data_are_preserved(self):
        model = valid_model()
        model["capabilities"] = [{"id": "CAP-1", "name": "CAPABILITY_SENTINEL", "description": "能力边界说明"}]
        model["requirements"][0]["capability_id"] = "CAP-1"
        model["data"] = [{"id": "DATA-1", "name": "DATA_SENTINEL", "fields": {"record_id": {"type": "string", "length": 32}}}]
        model["tailoring"] = [row for row in model["tailoring"] if row["clause"] != "3.5"]
        stream = json.dumps(generator.canonical_srs_to_sections(model, "review"), ensure_ascii=False)
        for token in ("CAPABILITY_SENTINEL", "DATA_SENTINEL", "record_id", "32"):
            self.assertIn(token, stream)

    def test_interface_requirements_stay_with_their_interface(self):
        model = valid_model()
        model["tailoring"] = [row for row in model["tailoring"] if row["clause"] != "3.3"]
        model["interfaces"] = [{"id": iid, "name": name, "csci_endpoint": "CSCI", "external_endpoint": name, "direction": "输入", "data": "受控消息", "protocol": "协议规范1"} for iid, name in (("IF-A", "接口甲"), ("IF-B", "接口乙"))]
        original = deepcopy(model["requirements"][0])
        model["requirements"] = [dict(original, id="REQ-A", clause="3.3", interface_id="IF-A"), dict(original, id="REQ-B", clause="3.3", interface_id="IF-B")]
        model["qualification"] = []
        model["forwardTrace"] = []
        model["reverseTrace"] = []
        sections = generator.canonical_srs_to_sections(model)
        a = next(row for row in sections if row["title"] == "接口甲（IF-A）")
        b = next(row for row in sections if row["title"] == "接口乙（IF-B）")
        self.assertIn("REQ-A", json.dumps(a))
        self.assertNotIn("REQ-B", json.dumps(a))
        self.assertIn("REQ-B", json.dumps(b))
        self.assertNotIn("REQ-A", json.dumps(b))

    def test_no_automatic_approved_tailoring_or_fake_trace(self):
        model = {"schema_version": "2.0", "metadata": {"document_type": "SRS", "title": "草稿"}, "requirements": []}
        stream = json.dumps(generator.canonical_srs_to_sections(model), ensure_ascii=False)
        self.assertNotIn("当前批准来源", stream)
        self.assertNotIn("本条无内容", stream)
        self.assertNotIn("来源到SRS正向追踪", stream)

    def test_parent_tailoring_is_explicitly_inherited(self):
        sections = generator.canonical_srs_to_sections(valid_model(), "review")
        row = next(row for row in sections if row["title"] == "计算机硬件需求")
        self.assertIn("继承 3.11", row["paragraphs"][0])

    def test_external_target_survives_generation(self):
        stream = json.dumps(generator.canonical_srs_to_sections(external_model(), "review"), ensure_ascii=False)
        self.assertIn("IRS引用：IRS-DOC / EXT-需求-7 / 3.2.1", stream)

    def test_low_level_review_bypass_and_mixed_content_are_rejected(self):
        demo = generator.srs_demo_content()
        with self.assertRaises(ValueError):
            generator._validate_lowlevel_sections(demo, "review")
        demo["requirements"] = valid_model()["requirements"]
        with self.assertRaises(ValueError):
            generator._validate_lowlevel_sections(demo, "draft")

    def test_failed_post_generation_check_preserves_existing_file(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "existing.docx"
            original = b"previous-approved-deliverable"
            output.write_bytes(original)
            with patch.object(generator, "require_page_number_structure", side_effect=ValueError("failed page check")):
                with self.assertRaises(ValueError):
                    generator.build_document(generator.SRS_TEMPLATE, valid_model(), output, "review")
            self.assertEqual(output.read_bytes(), original)
            self.assertEqual(sorted(p.name for p in Path(folder).iterdir()), ["existing.docx"])

    def test_actual_docx_preserves_minimal_requirement_and_data(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.docx"
            generator.build_document(generator.SRS_TEMPLATE, valid_model(), output, "review")
            document = Document(output)
            text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("REQ-任意#1", text)
            self.assertIn("软件应返回编号匹配的记录。", text)
            self.assertIn("所有需求具有同等优先级和关键性。", text)
            self.assertNotIn("当前批准来源", text)

    def test_existing_false_update_fields_is_enabled(self):
        document = Document(generator.SRS_TEMPLATE)
        flag = document.settings.element.find(qn("w:updateFields"))
        if flag is not None: flag.set(qn("w:val"), "false")
        generator.enable_update_fields(document)
        self.assertEqual(document.settings.element.find(qn("w:updateFields")).get(qn("w:val")), "true")

    def test_heading_uses_non_one_numbering_and_validates_levels(self):
        document = Document(generator.SRS_TEMPLATE)
        numbering = document.part.numbering_part.element
        numbering.find(qn("w:num")).set(qn("w:numId"), "42")
        generator.clear_body_content(document)
        generator.add_styled_paragraph(document, "自定义列表标题", "Heading 2")
        paragraph = document.paragraphs[-1]
        self.assertEqual(paragraph._p.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val")), "42")
        abstract = numbering.find(qn("w:abstractNum"))
        for row in list(abstract.findall(qn("w:lvl"))):
            if row.get(qn("w:ilvl")) == "2": abstract.remove(row)
        with self.assertRaises(ValueError):
            generator.add_styled_paragraph(document, "缺层级标题", "Heading 3")

    def test_direct_legacy_toc_cache_is_cleared_without_losing_sections(self):
        document = Document(generator.DEFAULT_TEMPLATE)
        sections = len(document.sections)
        generator.clear_body_content(document)
        generator.reset_toc_field(document)
        instructions = " ".join(node.text or "" for node in document._element.iter(qn("w:instrText")))
        self.assertIn("TOC", instructions)
        self.assertNotIn("PAGEREF _Toc", instructions)
        self.assertEqual(len(document.sections), sections)
        starts = {node.get(qn("w:id")) for node in document._element.iter(qn("w:bookmarkStart"))}
        ends = {node.get(qn("w:id")) for node in document._element.iter(qn("w:bookmarkEnd"))}
        self.assertEqual(starts, ends)

    def test_invalid_cli_input_returns_json_and_nonzero(self):
        with tempfile.TemporaryDirectory() as folder:
            content = Path(folder) / "invalid.json"
            content.write_text('{"metadata":[]}', encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPTS / "apply_gjb438c_template.py"), "--content-json", str(content), "--output", str(Path(folder) / "output.docx"), "--json"], capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(result.returncode, 2)
            report = json.loads(result.stdout)
            self.assertFalse(report["ok"])
            self.assertEqual(report["checks"]["visual_review"], "not_run")
            self.assertFalse((Path(folder) / "output.docx").exists())

    def test_review_cli_missing_fact_has_located_findings(self):
        with tempfile.TemporaryDirectory() as folder:
            model = valid_model()
            model["source_requirements"].append(dict(model["source_requirements"][0], id="SYS-MISSING"))
            content = Path(folder) / "model.json"
            content.write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPTS / "apply_gjb438c_template.py"), "--content-json", str(content), "--output", str(Path(folder) / "output.docx"), "--mode", "review", "--json"], capture_output=True, text=True, encoding="utf-8", env=dict(os.environ, PYTHONIOENCODING="cp1252"))
            self.assertEqual(result.returncode, 1)
            report = json.loads(result.stdout)
            self.assertTrue(any(row["rule_id"] == "SRS-SOURCE-UNCOVERED" and row["object_id"] == "SYS-MISSING" for row in report["issues"]))
            self.assertIn("来源", result.stdout)


if __name__ == "__main__":
    unittest.main()
