"""Observable authoring/profile behavior using synthetic public examples only."""
from copy import deepcopy
from pathlib import Path
import base64
import json
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_srs_model import valid_model, SCRIPTS
from srs_authoring import capability_outline
from srs_model import validate_srs_content
import apply_gjb438c_template as generator
from audit_gjb438c_docx import audit_docx
from docx import Document
from docx.oxml.ns import qn


def training_model():
    model = valid_model()
    model.pop("writing_profile")  # Exercise the new default.
    model["capability_overview"] = "记录查询平台由查询服务和导出服务组成。"
    model["expected_capability_roots"] = 2
    model["capabilities"] = [
        {"id": "SEARCH", "name": "查询服务", "order": 1},
        {"id": "QUERY", "name": "查询模块", "parent_id": "SEARCH"},
        {"id": "BY-ID", "name": "编号查询", "parent_id": "QUERY"},
        {"id": "EXPORT", "name": "导出服务", "order": 2},
    ]
    for node in model["capabilities"]:
        node.update(overview="按授权范围处理记录对象。", development_status="new", development_basis="示例合同第3条新研范围")
    req = model["requirements"][0]
    req.update(capability_id="BY-ID", exception="未知编号返回未找到，非法编号拒绝查询。", contract_trace={"status": "mapped", "source_ids": ["SYS-1"]})
    req2 = deepcopy(req)
    req2.update(id="REQ-EXPORT", statement="软件应导出编号匹配的记录。", capability_id="EXPORT")
    model["requirements"].append(req2)
    for group in ("qualification", "forwardTrace", "reverseTrace"):
        row = deepcopy(model[group][0]); row["requirement_id"] = req2["id"]; model[group].append(row)
    model["source_requirements"][0].update(kind="contract", confirmation="confirmed")
    model["figures"] = []
    for kind in ("composition", "use_case", "activity"):
        model["figures"].append({"id": "SYSTEM-" + kind, "type": kind, "title": "示例总体" + kind, "clause": "3.2", "reference_id": "SYS-DOC", "location": "图1"})
    for node in model["capabilities"]:
        for kind in ("composition", "activity"):
            ids = [r["id"] for r in model["requirements"] if r["capability_id"] == node["id"]]
            model["figures"].append({"id": node["id"] + "-" + kind, "type": kind, "title": node["name"] + kind, "capability_ids": [node["id"]], "requirement_ids": ids, "reference_id": "SYS-DOC", "location": "图2"})
    return model


class AuthoringTests(unittest.TestCase):
    def rules(self, model, mode="review", profile=None):
        return {r["rule_id"] for r in validate_srs_content(model, mode, profile) if r["severity"] == "error"}

    def test_default_training_and_explicit_standard_are_independent_of_mode(self):
        old = valid_model(); old.pop("writing_profile")
        self.assertIn("SRS-AUTHORING-OVERVIEW", self.rules(old))
        self.assertFalse(self.rules(old, "draft"))
        self.assertFalse(self.rules(old, profile="gjb-standard"))
        self.assertFalse(self.rules(training_model()))

    def test_unequal_depth_preserves_each_requirement_once(self):
        model = training_model()
        original = deepcopy(model)
        rows = capability_outline(model)
        self.assertEqual([(r["node"]["id"], r["level"], r["physical_clause"]) for r in rows], [("SEARCH", 3, "3.2.1"), ("QUERY", 4, "3.2.1.1"), ("BY-ID", 5, "3.2.1.1.1"), ("EXPORT", 3, "3.2.2")])
        sections = generator.canonical_srs_to_sections(model, "review")
        records = [table for section in sections for table in section.get("tables", []) if "需求陈述" in table["headers"]]
        self.assertEqual(len(records), 2)
        self.assertEqual(model, original)

    def test_bad_parent_cycle_duplicate_and_wrong_physical_clause_are_hard_errors(self):
        for change in (lambda m: m["capabilities"][1].update(parent_id="missing"), lambda m: m["capabilities"][0].update(parent_id="BY-ID"), lambda m: m["capabilities"].append(deepcopy(m["capabilities"][0])), lambda m: m["requirements"][0].update(physical_clause="3.2.2")):
            model = training_model(); change(model)
            self.assertTrue(self.rules(model, "draft"))

    def test_missing_leaf_overview_diagram_and_exception_remain_draft_findings(self):
        model = training_model()
        model["capabilities"].append({"id": "EMPTY", "name": "未细化模块"})
        model["requirements"][0].pop("exception")
        model["figures"] = [f for f in model["figures"] if f["id"] != "BY-ID-activity"]
        errors = self.rules(model)
        for rule in ("EMPTY-LEAF", "NODE-OVERVIEW", "NODE-DIAGRAM", "EXCEPTION", "ROOT-COUNT"):
            self.assertIn("SRS-AUTHORING-" + rule, errors)
        self.assertFalse(self.rules(model, "draft"))

    def test_contract_mapping_does_not_imply_approved_contract(self):
        model = training_model()
        model["source_requirements"][0]["confirmation"] = "pending"
        self.assertIn("SRS-AUTHORING-CONTRACT-BASELINE", self.rules(model))
        self.assertFalse(self.rules(model, "draft"))
        model["requirements"][0]["contract_trace"] = {"status": "pending", "basis": "保留来源需求，待合同范围确认", "tbd_ids": ["TBD-C"]}
        model["tbd"] = [{"id": "TBD-C", "issue": "合同范围确认", "owner": "项目责任人", "status": "open", "closure_condition": "提供确认范围", "affected_ids": [model["requirements"][0]["id"]]}]
        self.assertFalse(self.rules(model, "draft"))
        self.assertIn("SRS-AUTHORING-CONTRACT-PENDING", self.rules(model))
        model["requirements"][0]["contract_trace"]["tbd_ids"] = ["TBD-MISSING"]
        self.assertIn("SRS-AUTHORING-CONTRACT-TBD", self.rules(model, "draft"))

    def test_noncontract_source_cannot_be_claimed_as_contract_mapping(self):
        model = training_model(); model["source_requirements"][0]["kind"] = "plan"
        self.assertIn("SRS-AUTHORING-CONTRACT-KIND", self.rules(model))

    def test_word_depth_limit_is_explicit_without_flattening(self):
        model = valid_model()
        model["capabilities"] = [{"id": "N" + str(i), "name": "节点" + str(i), **({"parent_id": "N" + str(i - 1)} if i else {})} for i in range(7)]
        model["requirements"][0]["capability_id"] = "N6"
        self.assertEqual(capability_outline(model)[-1]["level"], 9)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "deep.docx"
            generator.build_document(generator.SRS_TEMPLATE, model, path, "review")
            doc = Document(path)
            self.assertEqual(next(p.style.name for p in doc.paragraphs if p.text == "节点6（N6）"), "Heading 9")
            instructions = " ".join(n.text or "" for n in doc._element.iter(qn("w:instrText")))
            self.assertIn('1-9', instructions)
        model["capabilities"].append({"id": "N7", "name": "超限节点", "parent_id": "N6"})
        self.assertIn("SRS-AUTHORING-TREE", self.rules(model, "draft"))

    def test_summary_figure_cannot_replace_module_activity_coverage(self):
        model = training_model()
        for fig in model["figures"]:
            if fig["id"] == "BY-ID-activity":
                fig["requirement_ids"] = []
        self.assertIn("SRS-AUTHORING-REQUIREMENT-DIAGRAM", self.rules(model))

    def test_tbd_context_is_local_and_tracks_file_identity_links(self):
        model = training_model()
        model["tbd"] = [{"id": "TBD-DEV", "issue": "研制范围", "owner": "示例角色", "status": "open", "closure_condition": "确认研制范围", "impact": "能力节点研制状态"}, {"id": "TBD-FILE", "issue": "来源文件身份", "owner": "示例角色", "status": "open", "closure_condition": "提供受控文件", "impact": "来源身份"}]
        for node in model["capabilities"]:
            node.update(development_status="tbd", development_basis="范围见TBD-DEV")
        model["references"][0].update(organization="待确认", tbd_ids=["TBD-FILE"])
        for row in model["forwardTrace"]:
            row["disposition"] = "映射已核，签署状态待确认"
        issues = validate_srs_content(model, "draft")
        self.assertNotIn("SRS-TBD-UNLINKED", {i["rule_id"] for i in issues})
        self.assertIn("SRS-TBD-OPEN", {i["rule_id"] for i in issues})
        model["requirements"][0]["statement"] = "软件应处理待确认的记录范围。"
        model["notes"] = "TBD-DEV和TBD-FILE均为开放事项。"
        issues = validate_srs_content(model, "draft")
        self.assertTrue(any(i["rule_id"] == "SRS-TBD-UNLINKED" and i["location"] == "/requirements/0/statement" for i in issues))
        model["references"][0]["tbd_ids"] = ["TBD-UNKNOWN"]
        self.assertIn("SRS-TBD-UNKNOWN", self.rules(model, "draft"))

    def test_closed_tbd_does_not_cover_still_unknown_development_status(self):
        model = training_model()
        model["capabilities"][0].update(development_status="tbd", development_basis="见TBD-DEV")
        model["tbd"] = [{"id": "TBD-DEV", "issue": "研制范围", "owner": "示例角色", "status": "closed", "closure_condition": "确认范围", "closure_evidence": "示例受控记录", "impact": "研制状态"}]
        self.assertTrue(any(i["rule_id"] == "SRS-TBD-UNLINKED" and i["location"] == "/capabilities/0/development_status" for i in validate_srs_content(model, "draft")))

    def test_saved_function_list_tampering_is_detected(self):
        with tempfile.TemporaryDirectory() as folder:
            model = training_model(); path = Path(folder) / "function.docx"
            generator.build_document(generator.SRS_TEMPLATE, model, path, "review")
            doc = Document(path)
            table = next(t for t in doc.tables if t.rows[0].cells[0].text == "唯一标识")
            table.rows[1].cells[0].text = "UNRELATED"
            doc.save(path)
            report = audit_docx(path, document_type="SRS", mode="draft", content=model)
            self.assertIn("SRS.MODEL_FUNCTION_LIST", {f["rule_id"] for f in report.findings})

    def test_saved_docx_keeps_hierarchy_and_detects_missing_figure(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            png = folder / "synthetic.png"
            png.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="))
            model = training_model()
            for fig in model["figures"]:
                fig.pop("reference_id"); fig.pop("location"); fig["path"] = str(png)
            path = folder / "review.docx"
            generator.build_document(generator.SRS_TEMPLATE, model, path, "review")
            doc = Document(path)
            node_heading = next(p for p in doc.paragraphs if p.text == "编号查询（BY-ID）")
            self.assertEqual(node_heading.style.name, "Heading 5")
            self.assertEqual(node_heading._p.find("./" + qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:ilvl")).get(qn("w:val")), "4")
            report = audit_docx(path, document_type="SRS", mode="review", content=model)
            self.assertTrue(report.ok, report.errors)
            self.assertEqual(report.checks["semantic_review"], "not_run")
            self.assertEqual(report.checks["visual_review"], "not_run")
            drawing = next(doc._element.iter(qn("w:drawing")))
            drawing.getparent().remove(drawing); doc.save(path)
            report = audit_docx(path, document_type="SRS", mode="draft", content=model)
            self.assertIn("SRS.MODEL_FIGURE", {f["rule_id"] for f in report.findings})

    def test_cli_can_explicitly_select_standard(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory); model = valid_model(); model.pop("writing_profile")
            source = folder / "input.json"; source.write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPTS / "apply_gjb438c_template.py"), "--content-json", str(source), "--output", str(folder / "out.docx"), "--writing-profile", "gjb-standard", "--mode", "review", "--json"], capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(json.loads(result.stdout)["writing_profile"], "gjb-standard")


if __name__ == "__main__":
    unittest.main()
